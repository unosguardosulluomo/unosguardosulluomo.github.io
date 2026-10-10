from pathlib import Path
import sys

from docx import Document


def replace_exact(document: Document, old: str, new: str) -> None:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == old]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one match, found {len(matches)} for: {old[:80]!r}")
    matches[0].text = new


def move_after(document: Document, text_to_move: str, destination_text: str) -> None:
    moving = [paragraph for paragraph in document.paragraphs if paragraph.text == text_to_move]
    destination = [paragraph for paragraph in document.paragraphs if paragraph.text == destination_text]
    if len(moving) != 1 or len(destination) != 1:
        raise RuntimeError(
            f"Unable to move paragraph: moving={len(moving)}, destination={len(destination)}"
        )
    destination[0]._p.addnext(moving[0]._p)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: revise_sindacato_vampiro_integrale.py INPUT.docx OUTPUT.docx")

    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)

    replacements = {
        "IL SINDACATO VAMPIRO":
            "LA POVERTÀ CONTAGIOSA",
        "La povertà che non vediamo, le famiglie che la sostengono e il mezzo miliardo di euro trattenuto ogni anno dalle pensioni":
            "Quando la povertà di uno finisce per impoverire anche chi gli sta più vicino",
        "INCHIESTA | SOCIETÀ | ITALIA":
            "GEOPOLITICA | SOCIETÀ | INCHIESTA",
        "In Italia più di 5,7 milioni di persone vivono in povertà assoluta. Sono oltre 2,2 milioni di famiglie. Non sono cifre con cui confezionare una notizia di passaggio: sono persone che, secondo la misurazione ISTAT riferita al 2024, non dispongono di quanto serve per sostenere una spesa essenziale adeguata alle proprie condizioni. Chi non riesce ad acquistare ciò di cui ha bisogno non può rimandare per sempre: rinuncia, si indebita, cerca aiuto.":
            "In Italia più di 5,7 milioni di persone vivono in povertà assoluta. Sono oltre 2,2 milioni di famiglie. Non sono cifre con cui confezionare una notizia di passaggio: sono persone che, secondo la misurazione dell’Istituto nazionale di statistica (ISTAT) riferita al 2024, non dispongono di quanto serve per sostenere una spesa essenziale adeguata alle proprie condizioni. Chi non riesce ad acquistare ciò di cui ha bisogno non può rimandare per sempre: rinuncia, si indebita, cerca aiuto.",
        "Il punto di partenza di questa inchiesta non è il sindacato. È una domanda molto più elementare: come vivono quelle persone? Chi mette sul tavolo il cibo quando il reddito non basta? Chi paga le scarpe del bambino, una bolletta in scadenza, le medicine, una spesa che non può più essere rinviata? E soprattutto: che cosa succede al bilancio di chi aiuta?":
            "Il punto di partenza di questa inchiesta è una domanda elementare: come vivono quelle persone? Chi mette sul tavolo il cibo quando il reddito non basta? Chi paga le scarpe del bambino, una bolletta in scadenza, le medicine, una spesa che non può più essere rinviata? E soprattutto: che cosa succede al bilancio di chi aiuta?",
        "Caritas: il volto delle persone che chiedono aiuto":
            "Caritas incontra e accompagna le persone",
        "Il Rapporto Caritas relativo al 2024 registra 277.775 persone accompagnate nei servizi censiti, attraverso 3.341 strutture informatizzate. È un numero importante, ma non rappresenta l’intera povertà italiana: fotografa una rete di servizi e di persone che hanno avuto accesso a quella rete, non tutti coloro che hanno bisogno di aiuto. Un assistito registrato può inoltre essere il riferimento di un intero nucleo familiare, con bisogni che investono più persone.":
            "Il Rapporto Caritas relativo al 2024 registra 277.775 persone accompagnate nei servizi censiti, attraverso 3.341 strutture informatizzate. Caritas incontra le persone, ne ascolta i bisogni e le accompagna attraverso centri e servizi che non si limitano all’aiuto alimentare. Il dato fotografa chi è entrato in quella rete e non rappresenta l’intera povertà italiana. Un assistito registrato può inoltre essere il riferimento di un intero nucleo familiare, con bisogni che investono più persone.",
        "Banco Alimentare: un’altra grande rete, non un’altra popolazione da sommare":
            "Banco Alimentare rifornisce gli enti che distribuiscono il cibo",
        "La Fondazione Banco Alimentare e la rete di organizzazioni territoriali che distribuiscono gli alimenti costituiscono un altro elemento fondamentale. Un dato di riferimento del 2023 parla di circa 1.794.000 persone raggiunte attraverso la rete di enti convenzionati. Non è il numero delle persone che ogni mattina si presentano direttamente a un magazzino del Banco: il Banco recupera e distribuisce derrate attraverso realtà locali che le consegnano alle persone e alle famiglie.":
            "La Fondazione Banco Alimentare opera su un livello diverso. Recupera le eccedenze e distribuisce gli alimenti agli enti territoriali convenzionati, che a loro volta li consegnano alle persone e alle famiglie. Il dato di riferimento del 2023 parla di circa 1.794.000 persone raggiunte attraverso questi enti. Non sono persone che si presentano direttamente ai magazzini del Banco e non costituiscono automaticamente una popolazione separata da quella incontrata da Caritas.",
        "Fra queste realtà possono comparire anche Caritas, associazioni di volontariato, realtà parrocchiali, Croce Rossa e organizzazioni assistenziali. Gli stessi nuclei possono ricevere servizi da più reti; dunque sarebbe sbagliato sommare 1,794 milioni ai 277.775 registrati dalla Caritas e proclamare un totale nazionale di assistiti senza sovrapposizioni.":
            "Fra gli enti riforniti possono comparire strutture Caritas, associazioni di volontariato, realtà parrocchiali, Croce Rossa e altre organizzazioni assistenziali. Banco Alimentare descrive quindi una filiera di approvvigionamento; Caritas descrive una rete che incontra e accompagna le persone. Le due funzioni possono incrociarsi e gli stessi nuclei possono ricevere più forme di assistenza. Sarebbe perciò sbagliato sommare 1,794 milioni ai 277.775 registrati dalla Caritas e presentarli come un totale nazionale privo di sovrapposizioni.",
        "Nella nostra discussione abbiamo eseguito un confronto elementare: 5.700.000 meno 1.794.000 fa circa 3.906.000. Il risultato aritmetico è corretto. Sarebbe invece falsa la conclusione che 3,9 milioni di poveri non ricevano alcun aiuto: alcuni sono sostenuti da altri enti, alcuni da servizi pubblici, alcuni da reti di vicinato, molti dai parenti. I due numeri appartengono inoltre a rilevazioni e anni differenti e non descrivono lo stesso universo statistico.":
            "Nella nostra discussione abbiamo eseguito un confronto elementare: 5.700.000 meno 1.794.000 fa circa 3.906.000. Il risultato aritmetico è corretto, ma non misura le persone rimaste senza aiuto. Il primo dato è una stima della povertà assoluta; il secondo conta persone raggiunte tramite enti convenzionati, fra i quali possono esserci anche strutture Caritas. Le rilevazioni riguardano anni e universi differenti. Alcune persone ricevono inoltre sostegno da servizi pubblici, altre associazioni, reti di vicinato o familiari.",
        "L’Osservatorio statistico INPS delle pensioni vigenti al 1° gennaio 2026 indica 2.822.563 prestazioni di vecchiaia inferiori a 750 euro mensili. Fra esse figurano numerose pensioni integrate e prestazioni di importo molto ridotto. Il dato, lo ripetiamo, riguarda gli assegni, non altrettanti pensionati che vivono esclusivamente con quella somma. Una persona può ricevere più trattamenti, una pensione ai superstiti oppure un reddito ulteriore.":
            "L’Osservatorio statistico dell’Istituto nazionale della previdenza sociale (INPS) sulle pensioni vigenti al 1° gennaio 2026 indica 2.822.563 prestazioni di vecchiaia inferiori a 750 euro mensili. Fra esse figurano numerose pensioni integrate e prestazioni di importo molto ridotto. Il dato, lo ripetiamo, riguarda gli assegni, non altrettanti pensionati che vivono esclusivamente con quella somma. Una persona può ricevere più trattamenti, una pensione ai superstiti oppure un reddito ulteriore.",
        "Non confondiamo inoltre l’attività associativa con i servizi di patronato e CAF. Il patronato svolge funzioni previste e finanziate anche da meccanismi pubblici: una parte dell’assistenza previdenziale istituzionale è accessibile anche ai non iscritti. Alcune prestazioni dei CAF possono essere a pagamento, anche per gli iscritti, altre gratuite secondo condizioni e convenzioni. Non sarebbe corretto dire che tutti i servizi siano sempre pagati a parte; sarebbe altrettanto scorretto sostenere che una tessera sindacale sia indispensabile per presentare domanda di pensione.":
            "Non confondiamo inoltre l’attività associativa con i servizi di patronato e con quelli dei centri di assistenza fiscale (CAF). Il patronato svolge funzioni previste e finanziate anche da meccanismi pubblici: una parte dell’assistenza previdenziale istituzionale è accessibile anche ai non iscritti. Alcune prestazioni dei CAF possono essere a pagamento, anche per gli iscritti, altre gratuite secondo condizioni e convenzioni. Non sarebbe corretto dire che tutti i servizi siano sempre pagati a parte; sarebbe altrettanto scorretto sostenere che una tessera sindacale sia indispensabile per presentare domanda di pensione.",
        "Sono state considerate ipotesi di investimento in BOT, un fondo presso INPS e un possibile coinvolgimento di Cassa Depositi e Prestiti. La distinzione è importante: l’INPS eroga le pensioni, Poste Italiane è uno degli intermediari attraverso cui i pensionati ricevono i pagamenti, e CDP gestisce attività finanziarie legate anche al risparmio postale. Non si può dire che tutto il flusso pensionistico italiano sia gestito da Poste, né che i tre soggetti siano la stessa cassa.":
            "Sono state considerate ipotesi di investimento in Buoni ordinari del Tesoro (BOT), un fondo presso l’INPS e un possibile coinvolgimento di Cassa depositi e prestiti (CDP). La distinzione è importante: l’INPS eroga le pensioni, Poste Italiane è uno degli intermediari attraverso cui i pensionati ricevono i pagamenti e CDP gestisce attività finanziarie legate anche al risparmio postale. Non si può dire che tutto il flusso pensionistico italiano sia gestito da Poste, né che i tre soggetti siano la stessa cassa.",
        "Sottotitolo: La povertà che non vediamo, le famiglie che la sostengono e il mezzo miliardo di euro trattenuto ogni anno dalle pensioni":
            "Sottotitolo: Quando la povertà di uno finisce per impoverire anche chi gli sta più vicino",
        "Slug permanente proposto: il-sindacato-vampiro":
            "Slug permanente proposto: article-geopolitica-societa-poverta-contagiosa.html",
        "URL canonico proposto: https://unosguardosulluomo.github.io/article-il-sindacato-vampiro.html":
            "URL canonico proposto: https://unosguardosulluomo.github.io/article-geopolitica-societa-poverta-contagiosa.html",
        "Neppure il titolo autorizza equivoci: «Il sindacato vampiro» è una formula critica e satirica sul flusso di risorse e sull’efficacia della rappresentanza, non un’accusa penale. Le organizzazioni ricevono denaro attraverso un meccanismo normativo esistente; proprio quel meccanismo è l’oggetto della contestazione editoriale.":
            "Il titolo «La povertà contagiosa» descrive un meccanismo economico e familiare, non attribuisce una colpa alle persone povere. Quando il reddito non basta, il costo si sposta su figli, genitori, fratelli e nonni; il flusso delle trattenute sindacali è una delle risorse sulle quali il dossier interroga le priorità collettive.",
        "Il titolo di questa inchiesta è «IL SINDACATO VAMPIRO». Il suo soggetto, però, è l’uomo. Quello che chiede aiuto, quello che lo offre a costo di privarsi del necessario e quello che, arrivato alla vecchiaia, avrebbe diritto almeno a non dover scegliere ogni giorno che cosa sacrificare.":
            "La povertà diventa contagiosa quando il bisogno di una persona si trasferisce su chi le sta accanto. Il soggetto di questa inchiesta è l’uomo: quello che chiede aiuto, quello che lo offre a costo di privarsi del necessario e quello che, arrivato alla vecchiaia, avrebbe diritto almeno a non dover scegliere ogni giorno che cosa sacrificare.",
        "Titolo: IL SINDACATO VAMPIRO":
            "Titolo: LA POVERTÀ CONTAGIOSA",
    }

    for old, new in replacements.items():
        replace_exact(document, old, new)

    move_after(
        document,
        "È da qui che nasce la domanda successiva: esistono flussi di denaro ricorrenti, già prelevati dai redditi pensionistici, ai quali il Paese potrebbe attribuire una destinazione sociale diversa?",
        "Non vogliamo organizzare una gara della miseria tra 620 e 546 euro. Vogliamo mostrare quanto sia ridotto il margine economico di partenza e quanto conti anche una mensilità aggiuntiva.",
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"saved: {output}")
    print(f"replacements: {len(replacements)}")


if __name__ == "__main__":
    main()
