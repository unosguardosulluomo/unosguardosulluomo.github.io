from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\utente\OneDrive\Documenti\ChatGPT\sito uno sguardo sull'uomo")
OUT = ROOT / "DOSSIER_SIHER_LA_CAUTELA_ESCLUSA.docx"
HERO = Path(r"C:\Users\utente\.codex\generated_images\01a19274-9d04-7123-b79c-a64d137d1956\exec-196719eb-8934-414d-a662-aba98fafc846.png")
PROFILE = Path(r"C:\Users\utente\.codex\generated_images\01a19274-9d04-7123-b79c-a64d137d1956\exec-0d95c86c-c3c7-418a-970f-08c0e131b053.png")


def set_font(run, name="Aptos", size=11, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color="555555")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, name="Aptos Display", size=16 if level == 1 else 12, bold=True)
    return p


def add_body(doc, text, bold_lead=None, keep=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_quote(doc, text, attribution=None):
    p = doc.add_paragraph(style="Quote")
    p.paragraph_format.left_indent = Inches(0.42)
    p.paragraph_format.right_indent = Inches(0.30)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    r = p.add_run(f"“{text}”")
    set_font(r, name="Aptos", size=10.5, italic=True)
    if attribution:
        r2 = p.add_run(f"\n{attribution}")
        set_font(r2, name="Aptos", size=9.5, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)
    return p


def add_image(doc, path, alt_text, width=7.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(9)
    cap.paragraph_format.keep_with_next = True
    cr = cap.add_run("Immagine prodotta da Uno Sguardo sull’Uomo")
    set_font(cr, size=8.5, italic=True, color="555555")


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.58)
sec.left_margin = Inches(0.70)
sec.right_margin = Inches(0.70)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True

for style_name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12)):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = style_name != "Title"
    st.font.color.rgb = RGBColor(0, 0, 0)

# Remove any inherited paragraph border from Word's built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

if "Deck" not in styles:
    deck_style = styles.add_style("Deck", WD_STYLE_TYPE.PARAGRAPH)
else:
    deck_style = styles["Deck"]
deck_style.font.name = "Aptos"
deck_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
deck_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
deck_style.font.size = Pt(12.5)
deck_style.font.color.rgb = RGBColor(0, 0, 0)
deck_style.paragraph_format.space_after = Pt(8)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("UNO SGUARDO SULL’UOMO")
set_font(r, name="Aptos", size=10.5, bold=True)

add_image(
    doc,
    HERO,
    "Composizione editoriale con la gip Tiziana Landoni, fascicoli e analisi digitale, e Youssef Siher presso un’uscita",
    width=7.05,
)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_before = Pt(2)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run("Caso Siher la cautela esclusa")
set_font(tr, name="Aptos Display", size=25, bold=True)

deck = doc.add_paragraph(style="Deck")
dr = deck.add_run(
    "La scarcerazione decisa dalla gip Tiziana Landoni, il materiale digitale ancora da esaminare e le minacce emerse nelle chat del Movimento 17 ottobre"
)
set_font(dr, size=12.5)

meta = doc.add_paragraph()
meta.paragraph_format.space_before = Pt(2)
meta.paragraph_format.space_after = Pt(0)
mr = meta.add_run("DOSSIER  |  Chiusura redazionale 31 ottobre 2026  |  Tempo di lettura 18 minuti")
set_font(mr, size=8.5, bold=True, color="444444")

meta2 = doc.add_paragraph()
meta2.paragraph_format.space_after = Pt(0)
mr2 = meta2.add_run("Macroargomento Geopolitica  |  Microargomento Politica italiana")
set_font(mr2, size=8.5, color="444444")

doc.add_page_break()

# Footer from page 2 onward (same section, cover intentionally shares footer but first-page different).
sec.different_first_page_header_footer = True
footer = sec.footer
fp = footer.paragraphs[0]
fp.text = ""
ft = footer.add_table(rows=1, cols=2, width=Inches(7.1))
ft.autofit = False
set_repeat_table_header(ft.rows[0])
ft.columns[0].width = Inches(6.1)
ft.columns[1].width = Inches(1.0)
left = ft.rows[0].cells[0]
right = ft.rows[0].cells[1]
left.width = Inches(6.1)
right.width = Inches(1.0)
set_cell_margins(left, top=0, start=0, bottom=0, end=0)
set_cell_margins(right, top=0, start=0, bottom=0, end=0)
lp = left.paragraphs[0]
lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
lr_footer = lp.add_run("UNO SGUARDO SULL’UOMO  •  Caso Siher la cautela esclusa")
set_font(lr_footer, size=8.2, color="555555")
rp = right.paragraphs[0]
rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_field(rp, "PAGE")
tbl_pr = ft._tbl.tblPr
tbl_borders = OxmlElement("w:tblBorders")
for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "nil")
    tbl_borders.append(el)
tbl_pr.append(tbl_borders)

# Opening
lead = doc.add_paragraph()
lead.paragraph_format.space_after = Pt(10)
lead.paragraph_format.keep_together = True
lr = lead.add_run(
    "La decisione della giudice per le indagini preliminari Tiziana Landoni ha escluso ogni forma di cautela nei confronti di Youssef Siher. Il punto non è pretendere una condanna anticipata. È capire se fosse ragionevole liberarlo senza alcuna misura mentre gli accertamenti digitali non erano conclusi e mancavano due giorni a una nuova manifestazione nazionale."
)
set_font(lr, name="Aptos", size=13, bold=True)

add_body(
    doc,
    "Il 28 ottobre 2026 la giudice per le indagini preliminari (GIP) di Milano non ha convalidato il fermo disposto dalla Procura e ha respinto la richiesta di custodia cautelare in carcere. Siher è tornato libero. La decisione non ha chiuso l’indagine e non equivale a un’assoluzione, ma ha prodotto un effetto immediato: nessun vincolo personale, proprio nella fase in cui telefono e computer potevano ancora restituire elementi nuovi."
)
add_body(
    doc,
    "Landoni non ha definito la vicenda una bravata. Tuttavia, leggendo i contenuti direttamente attribuiti a Siher come disobbedienza civile e libera espressione, e negando ogni misura, il provvedimento ha adottato l’interpretazione meno allarmante del rischio prima che il quadro digitale fosse completo. È questa scelta, non l’indipendenza della magistrata né la legittimità del controllo giudiziario sul fermo, il centro dell’indagine."
)

add_heading(doc, "Un adulto con capacità organizzative", 1)
add_body(
    doc,
    "Youssef Siher ha trent’anni. Le cronache lo descrivono come ricercatore di geopolitica, studioso dei movimenti politici islamici e giornalista freelance. Non emerge il profilo di un adolescente trascinato occasionalmente dalla piazza, ma quello di un adulto capace di costruire e amministrare strumenti di comunicazione con migliaia di destinatari."
)
add_body(
    doc,
    "Secondo il Corriere della Sera, il canale Telegram del Movimento 17 ottobre aveva circa 11.000 iscritti e la pagina Instagram circa 39.000. I canali sarebbero stati creati la sera del 18 ottobre, dopo le devastazioni del 17 e 18 ottobre a Milano. La loro nascita successiva impedisce di attribuire automaticamente a Siher l’organizzazione di quelle violenze, ma rende evidente la rapidità con cui seppe inserirsi nella mobilitazione e trasformarla in una rete digitale ampia."
)
add_body(
    doc,
    "Siher si era sposato nel settembre 2026. Il matrimonio, da solo, non dimostra responsabilità o pericolosità. In questa vicenda, però, la dimensione familiare si intreccia direttamente con quella organizzativa: secondo l’Agenzia Nazionale Stampa Associata (ANSA), la moglie era indagata nello stesso procedimento e i due coniugi erano gli amministratori dei canali; nelle aree di pubblicazione riservate agli amministratori sarebbero stati gli unici autorizzati a scrivere. Questo dato non consente di attribuire loro ogni commento comparso negli spazi di discussione collegati, ma rende ancora meno credibile la rappresentazione di un’attività estemporanea o ingenua."
)

add_heading(doc, "La sequenza dei fatti", 1)
timeline = [
    ("17 e 18 ottobre 2026", "A Milano le manifestazioni degenerano in saccheggi, devastazioni e scontri. Il numero degli arresti riportato dalle cronache riguarda i fatti di piazza e non prova una responsabilità di Siher."),
    ("18 ottobre", "Nascono i canali del Movimento 17 ottobre, secondo la ricostruzione del Corriere della Sera."),
    ("24 ottobre", "Siher viene individuato e portato in Questura. Nella notte tra il 24 e il 25 ottobre la Procura dispone il fermo."),
    ("25 ottobre", "Viene trasferito a San Vittore. L’indagine coinvolge anche la moglie, indicata come coamministratrice."),
    ("28 ottobre", "Landoni non convalida il fermo, respinge la custodia cautelare e dispone la liberazione."),
    ("29 e 30 ottobre", "Diventano pubbliche ulteriori chat, comprese le frasi contro Giovanni Donzelli. Numerosi consiglieri del Consiglio superiore della magistratura chiedono una pratica a tutela di Landoni e il Tribunale di Milano chiede di valutarne la protezione per le minacce ricevute."),
    ("30 ottobre", "Si svolge lo sciopero generale con cortei di studenti e lavoratori. A Milano partecipano migliaia di persone; nei cortei compaiono anche messaggi di solidarietà a Siher."),
    ("31 ottobre", "Il Corriere della Sera riferisce che la Procura valuta con prudenza l’impugnazione e potrebbe attendere nuovi elementi dagli accertamenti digitali."),
]
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(1.25)
table.columns[1].width = Inches(5.65)
table.rows[0].cells[0].width = Inches(1.25)
table.rows[0].cells[1].width = Inches(5.65)
hdr = table.rows[0].cells
hdr[0].text = "Data"
hdr[1].text = "Evento"
for c in hdr:
    shade_cell(c, "243746")
    set_cell_margins(c)
    for p in c.paragraphs:
        for run in p.runs:
            set_font(run, size=9.5, bold=True, color="FFFFFF")
set_repeat_table_header(table.rows[0])
for idx, (date, event) in enumerate(timeline):
    cells = table.add_row().cells
    cells[0].text = date
    cells[1].text = event
    if idx % 2:
        shade_cell(cells[0], "F3F6F8")
        shade_cell(cells[1], "F3F6F8")
    for c in cells:
        set_cell_margins(c)
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_font(run, size=9.2, bold=(c is cells[0]))
set_table_borders(table)
doc.add_paragraph().paragraph_format.space_after = Pt(0)

add_heading(doc, "Come Landoni ha letto i messaggi", 1)
add_body(
    doc,
    "Il provvedimento ha separato le parole direttamente attribuite a Siher dalle violenze già avvenute. Secondo i passaggi pubblicati da più testate, nei post e nel comunicato l’indagato non avrebbe mai prospettato il ricorso ad azioni violente né indicato le devastazioni del 17 e 18 ottobre come modello della mobilitazione successiva. La giudice ha inoltre escluso un concreto pericolo di fuga."
)
add_quote(
    doc,
    "Nei post e nel comunicato [...] l’indagato non prospetta mai il ricorso ad azioni violente",
    "Passaggio dell’ordinanza riportato da Sky TG24, RaiNews e altre testate",
)
add_body(
    doc,
    "Questa lettura si concentra sull’autore materiale delle parole e sul loro significato letterale. La domanda editoriale è se un amministratore di canali creati per mobilitare migliaia di persone possa essere valutato soltanto cercando un ordine esplicito firmato da lui. Un’infrastruttura digitale può orientare tempi, luoghi, obiettivi e clima della protesta anche senza formulare direttamente un comando violento. La responsabilità penale resta individuale; la valutazione cautelare, però, riguarda anche il pericolo concreto che un ruolo organizzativo possa produrre."
)
add_body(
    doc,
    "Il silenzio dell’amministratore davanti a un messaggio violento non equivale automaticamente ad adesione. Occorre conoscere chi poteva cancellare, quali contenuti furono rimossi, quanto rimasero visibili, quali reazioni ricevettero e se Siher intervenne. Proprio perché questi dati richiedevano un esame tecnico, la rapidità con cui è stata esclusa ogni cautela diventa la questione decisiva."
)

add_heading(doc, "Nessuna misura mentre gli accertamenti continuavano", 1)
add_body(
    doc,
    "La convalida del fermo e l’applicazione di una misura cautelare sono decisioni distinte. L’articolo 391 del Codice di procedura penale prevede che, quando sussistono le condizioni dell’articolo 273 e una delle esigenze indicate dall’articolo 274, il giudice possa applicare una misura coercitiva. L’articolo 299 impone poi di revocarla o sostituirla quando quelle condizioni vengono meno o si attenuano."
)
add_body(
    doc,
    "La legge non permette di trattenere una persona soltanto per concedere più tempo agli investigatori. Permette però una misura temporanea e revocabile quando il giudice riconosce gravi indizi e concrete esigenze cautelari. Gli strumenti, dunque, esistevano. Landoni ha ritenuto che ne mancassero i presupposti e non ha applicato alcun vincolo. La critica riguarda precisamente questa valutazione: davvero ruolo, capacità di mobilitazione, contesto e imminenza di un’altra manifestazione non giustificavano nessuna cautela proporzionata?"
)
add_body(
    doc,
    "Le fonti pubbliche documentano la richiesta della Procura di applicare la custodia in carcere, respinta dalla GIP. Non consentono di ricostruire integralmente se e come siano state discusse tutte le alternative meno afflittive. Il dato certo resta l’esito: Siher è uscito senza alcuna misura personale."
)

add_image(
    doc,
    PROFILE,
    "Composizione editoriale di Youssef Siher a una scrivania con dispositivi digitali e una rete di comunicazione",
    width=7.0,
)

add_heading(doc, "Due giorni prima della nuova piazza nazionale", 1)
add_body(
    doc,
    "La liberazione è avvenuta il 28 ottobre. Il 30 ottobre era già previsto lo sciopero generale con mobilitazioni studentesche e cortei nazionali. A Milano le autorità attendevano fino a diecimila manifestanti e mantenevano alta l’attenzione dopo le violenze del fine settimana precedente. La manifestazione si è poi svolta con migliaia di studenti e lavoratori; furono esposti anche messaggi di solidarietà a Siher."
)
add_body(
    doc,
    "Il calendario non prova che Siher avrebbe commesso nuovi reati. Incide però sulla prudenza richiesta nella valutazione del rischio. La GIP ha escluso ogni misura prima della nuova mobilitazione, mentre il materiale sequestrato poteva ancora chiarire la rete dei contatti, la programmazione delle comunicazioni e la gestione effettiva dei canali. Una cautela temporanea non sarebbe stata una condanna: sarebbe stata una risposta revocabile a un quadro ancora in formazione, se il giudice avesse riconosciuto i presupposti di legge."
)

add_heading(doc, "Le chat contro Giovanni Donzelli", 1)
add_body(
    doc,
    "Il 30 ottobre il Giornale ha pubblicato screenshot attribuiti agli spazi Telegram collegati al Movimento 17 ottobre. Secondo l’articolo, quei messaggi non erano confluiti nella richiesta iniziale di fermo né nel provvedimento di liberazione. Non disponiamo dell’esportazione completa della conversazione e non attribuiamo le due frasi a Siher. Il loro contenuto, però, è esplicito e individua un bersaglio politico per nome."
)
add_quote(doc, "Andate sotto il Parlamento a migliaia… E magari spaccare il muso anche a Donzelli")
add_quote(doc, "Se qualcuno gli lancia qualcosa stile statuetta del Duomo a Berlusconi…")
add_body(
    doc,
    "Queste parole superano l’ostilità politica generica. Evocano un’aggressione fisica e richiamano l’attacco subito da Silvio Berlusconi nel 2009. Il punto da accertare non è soltanto chi le abbia scritte: occorre stabilire quando furono pubblicate, chi le vide, quanto rimasero online, chi disponeva dei poteri di moderazione e quali condotte seguirono."
)
add_body(
    doc,
    "Giovanni Donzelli non era un bersaglio casuale né privo di una storia di esposizione. Responsabile dell’organizzazione di Fratelli d’Italia (FdI), proveniente dalla destra universitaria e dal Fronte universitario d’azione nazionale (FUAN), è da tempo oggetto di contestazioni e minacce. Nel febbraio 2023 il Viminale dispose per lui una scorta nel contesto delle tensioni sul caso Cospito. Nel giugno 2026, a Firenze, comparve anche la scritta «Donzelli a testa in giù». Le nuove frasi colpiscono quindi una persona il cui rischio era già stato riconosciuto dallo Stato."
)

add_heading(doc, "Le reazioni politiche rimaste fuori dalla prima sintesi", 1)
add_body(
    doc,
    "Il materiale originariamente fornito alla redazione conteneva dichiarazioni politiche complete che l’archivio di lavoro aveva abbreviato. Nel dossier definitivo vengono recuperate perché documentano il modo in cui la vicenda è stata letta da Fratelli d’Italia. Sono dichiarazioni di parte e vanno considerate come tali."
)
add_quote(
    doc,
    "Rivolgo la piena solidarietà mia e dei deputati di Fratelli d’Italia a Giovanni Donzelli per le minacce emerse nelle chat di Siher. Colpisce, tuttavia, l’assordante silenzio della sinistra davanti a messaggi che inneggiano alla violenza e indicano esponenti di Fratelli d’Italia come bersagli. Ci chiediamo dove sono finiti gli appelli alla difesa della democrazia e alla condanna dell’odio politico. Ci chiediamo se i partiti d’opposizione scelgano di adottare consapevolmente questa precisa linea. Evitare attentamente di pestare i piedi ai movimenti che in questi giorni stanno mettendo a dura prova le Forze dell’ordine. Su minacce e intimidazioni non può esserci alcun conflitto d’interesse: per questo aspettiamo una presa di posizione chiara e senza ambiguità da parte di tutte le forze politiche.",
    "Galeazzo Bignami, dichiarazione riportata nel materiale fornito alla redazione",
)
add_quote(
    doc,
    "A quanto pare, tra le importanti rivendicazioni del Movimento 17 ottobre, oltre al diritto allo studio, al futuro dei giovani e alle tante questioni che animano le proteste studentesche, c’è anche il desiderio di spaccare la testa a Giovanni Donzelli. Almeno a giudicare da quello che si legge nelle chat degli organizzatori. Ci chiediamo se queste siano le nuove forme di partecipazione democratica che qualcuno continua a difendere e giustificare.",
    "Francesco Torselli, europarlamentare di Fratelli d’Italia",
)
add_quote(
    doc,
    "A Giovanni va la mia vicinanza. Manifestare, contestare e dissentire sono diritti sacrosanti, ma augurarsi aggressioni fisiche e trasformare chi la pensa diversamente in un bersaglio non ha nulla a che vedere con il diritto di protesta. È violenza, anche quando viene scritta comodamente dietro una tastiera. E allora vorremmo sapere se la sinistra abbia intenzione di prendere le distanze anche da queste minacce. Perché non si può condannare la violenza soltanto quando fa comodo.",
    "Francesco Torselli",
)
add_quote(
    doc,
    "Costoro, che si sentono evidentemente impuniti perché nascosti in un gruppo, ignorano che Giovanni non si farà certo intimorire da vigliaccate simili. Perché è rimasta la stessa persona che ai tempi dell’università, quando militava nel Fuan, non si lasciava spaventare da situazioni pericolose, a rischio e tantomeno dai violenti ideologici. Ci aspettiamo adesso una condanna unanime.",
    "Parlamentari toscani di Fratelli d’Italia, dichiarazione collettiva",
)

add_heading(doc, "La protezione di Landoni e la reazione delle istituzioni", 1)
add_body(
    doc,
    "Le minacce contro Landoni sono inaccettabili e devono essere valutate senza esitazioni. Il 29 ottobre i vertici del Tribunale di Milano hanno iniziato a monitorare i social. Il giorno successivo il presidente del Tribunale Fabio Roia ha scritto al prefetto e al questore, oltre che ai vertici della Procura, chiedendo di verificare se fossero necessarie misure di protezione. L’ANSA ha riferito che almeno due messaggi allegati alla lettera contenevano minacce esplicite."
)
add_body(
    doc,
    "Parallelamente, numerosi consiglieri del Consiglio superiore della magistratura (CSM) hanno chiesto l’apertura di una pratica a tutela della GIP dopo le dichiarazioni della presidente del Consiglio Giorgia Meloni, che aveva definito la decisione «vergognosa». La richiesta ha sostenuto che una decisione giudiziaria può essere criticata anche severamente, ma non deve essere presentata come ostacolo alla sicurezza o mortificazione delle forze dell’ordine."
)
add_body(
    doc,
    "La tutela della magistrata non è il problema. Il punto è la diversa velocità con cui le istituzioni riconoscono il rischio. Per Landoni, nuove minacce online hanno attivato monitoraggio, lettere formali e un’iniziativa consiliare. Le frasi contro Donzelli, già protetto per minacce pregresse, sono invece emerse dopo la scarcerazione e, secondo la ricostruzione giornalistica, non avevano pesato sul fascicolo iniziale. Il confronto non chiede di ridurre la protezione della giudice: chiede perché il nuovo materiale relativo a un bersaglio già esposto non abbia imposto una rivalutazione altrettanto visibile e tempestiva della posizione di Siher e della gestione dei suoi spazi digitali."
)

add_heading(doc, "La Procura dopo il rigetto", 1)
add_body(
    doc,
    "Il 31 ottobre il Corriere della Sera ha descritto come «più no che sì» l’orientamento della Procura sul ricorso, qualora il quadro degli elementi fosse rimasto invariato. L’articolo collega la prudenza al rischio di un secondo rigetto e alla possibilità che telefono e computer producano nuovi riscontri. La Procura può proseguire l’indagine a piede libero e contestare la valutazione giuridica già compiuta; non è obbligata a disporre di prove nuove per impugnare ogni profilo, ma nuovi elementi renderebbero più solida l’iniziativa cautelare."
)
add_body(
    doc,
    "Non è una battaglia contro la Procura. È l’effetto strategico della decisione di Landoni. Gli elementi iniziali non sono stati cancellati, ma hanno ricevuto un giudizio negativo nella fase cautelare. Ripresentarli senza un rafforzamento espone l’accusa a una seconda pronuncia sfavorevole. Attendere gli accertamenti digitali può quindi essere la scelta più prudente per l’indagine, ma lascia aperta la domanda sulla prudenza adottata nei giorni immediatamente precedenti alla manifestazione del 30 ottobre."
)

add_heading(doc, "Una decisione formalmente motivata ma poco prudente", 1)
add_body(
    doc,
    "Landoni ha esercitato il compito che la legge assegna alla GIP: controllare il fermo e impedire che la custodia diventi automatica. La sua ordinanza ha distinto la protesta politica dall’istigazione e ha ritenuto insufficienti gli elementi presentati. Questa è la motivazione giuridica. La valutazione pubblica può però interrogare il perimetro scelto dal provvedimento."
)
add_body(
    doc,
    "Siher era un trentenne con esperienza di ricerca, comunicazione e organizzazione; insieme alla moglie amministrava strumenti capaci di raggiungere decine di migliaia di persone. Le manifestazioni precedenti erano degenerate. Una nuova giornata nazionale era imminente. Il materiale digitale non era esaurito. In quel contesto, l’assenza di un ordine violento scritto personalmente da Siher bastava davvero a escludere ogni pericolo cautelare?"
)
add_body(
    doc,
    "Il fermo non è una condanna e una misura cautelare non deve diventare una pena anticipata. Ma anche la libertà senza condizioni è una scelta, non un esito neutro. Fra la custodia cautelare in carcere e l’assenza di qualsiasi vincolo esistono misure graduate, soggette a controllo e revocabili quando i presupposti vengono meno. Le fonti pubbliche non permettono di sapere quali alternative fossero concretamente percorribili nel fascicolo; permettono però di constatare che nessuna fu applicata."
)
add_body(
    doc,
    "La nostra conclusione è editoriale e resta distinta da un accertamento penale: la GIP ha privilegiato una lettura individuale e letterale dei messaggi, sottovalutando il peso organizzativo dell’amministratore, il contesto della rete e il momento in cui decideva. Non ha trattato formalmente Siher come l’autore di una bravata; l’effetto pratico della decisione, tuttavia, è stato quello di escludere qualsiasi cautela davanti a un adulto strutturato e a un quadro ancora incompleto."
)

add_heading(doc, "Le domande che restano", 1)
for item in [
    "Quali messaggi erano effettivamente presenti nella richiesta della Procura e quali sono emersi soltanto dopo la liberazione?",
    "Chi scrisse le frasi contro Donzelli, quando furono pubblicate e quali interventi di moderazione seguirono?",
    "Quale ruolo ebbero Siher e la moglie nelle aree di pubblicazione, nei gruppi di discussione e nella programmazione delle convocazioni?",
    "Perché nessuna misura fu ritenuta proporzionata mentre gli accertamenti informatici erano ancora in corso?",
    "La Procura rinuncerà all’impugnazione perché condivide la debolezza del quadro cautelare o perché considera troppo alto il rischio di una seconda bocciatura?",
    "I nuovi elementi digitali porteranno a una rivalutazione della posizione di Siher o confermeranno la lettura adottata da Landoni?",
]:
    add_bullet(doc, item)

add_heading(doc, "Fonti", 1)
sources = [
    "ANSA, Amministratori chat Movimento 17 ottobre solo Siher e sua moglie, 25 ottobre 2026.",
    "ANSA, Siher torna in libertà Meloni attacca la gip, 28 ottobre 2026.",
    "ANSA, Il Csm chiede una pratica a tutela della gip del caso Siher, Luca Capponi, 29 ottobre 2026.",
    "ANSA, Caso Siher Tribunale chiede di verificare se serve tutela per gip, 30 ottobre 2026.",
    "ANSA, A Milano già centinaia in piazza cambia il percorso del corteo degli studenti, 30 ottobre 2026.",
    "Corriere della Sera Milano, Chi è Youssef Siher la mente dietro il Movimento 17 ottobre, 24 ottobre 2026.",
    "Corriere della Sera Milano, Scontri a Milano indagata anche la moglie di Youssef Siher, 25 ottobre 2026.",
    "Corriere della Sera Milano, Youssef Siher non convalida del fermo il ricorso possibile soltanto con nuovi elementi di indagine, Luigi Ferrarella, 31 ottobre 2026.",
    "Sky TG24, Proteste Milano fermato Siher, 25 ottobre 2026.",
    "Sky TG24, Youssef Siher il gip di Milano lo rimette in libertà, 28 ottobre 2026.",
    "RaiNews, Youssef Siher è in carcere indagata anche la moglie, 25 ottobre 2026.",
    "RaiNews, Cortei Youssef Siher è libero chiesta in Csm tutela della gip, 29 ottobre 2026.",
    "RaiNews, Disposta la scorta anche per Donzelli, aggiornamento del 3 febbraio 2023 sul caso Cospito.",
    "Il Giornale, Pure Donzelli nel mirino di Siher Ira FdI sulla gip, 30 ottobre 2026.",
    "Secolo d’Italia, Spaccate il muso a Donzelli il messaggio choc nelle chat di Siher, 30 ottobre 2026.",
    "Fratelli d’Italia, Siher con nuove chat si rivaluti posizione Solidarietà a Donzelli per minacce, 29 ottobre 2026.",
    "Normattiva, Codice di procedura penale, articoli 273, 274, 299, 310 e 391, testo vigente consultato il 31 ottobre 2026.",
    "Normattiva, Codice penale, articolo 414, testo vigente consultato il 31 ottobre 2026.",
]
for src in sources:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    r = p.add_run(src)
    set_font(r, size=9.3)

doc.core_properties.title = "Caso Siher la cautela esclusa"
doc.core_properties.subject = "Dossier sulla decisione cautelare nel caso Youssef Siher"
doc.core_properties.author = "Redazione Uno Sguardo sull’Uomo"
doc.core_properties.keywords = "Youssef Siher, Tiziana Landoni, Movimento 17 ottobre, Giovanni Donzelli, misure cautelari"
doc.core_properties.comments = "Versione definitiva Word chiusa il 31 ottobre 2026"

doc.save(OUT)
print(OUT)
