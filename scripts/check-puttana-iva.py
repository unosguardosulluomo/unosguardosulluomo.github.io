"""Validate the IVA dossier, metadata, assets, links and editorial placements."""

import html
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/puttana-iva.json").read_text(encoding="utf-8"))
page = (ROOT / D["slug"]).read_text(encoding="utf-8")


def normalize(value):
    return re.sub(r"\s+", " ", value).strip()


url = "https://unosguardosulluomo.github.io/" + D["slug"]
assert D["slug"] == "article-puttana-iva.html"
assert f'<link rel="canonical" href="{url}">' in page
assert f'<meta property="og:url" content="{url}">' in page
assert f'<meta property="article:published_time" content="{D["datePublished"]}">' in page
assert f'<meta property="article:modified_time" content="{D["dateModified"]}">' in page
assert f'data-macro-category="{D["macroCategory"]}"' in page
assert f'data-micro-category="{D["microCategory"]}"' in page
assert page.count('<figure class="') == 2
assert page.count("Immagine prodotta da Uno Sguardo sull’Uomo.") == 2
assert "intelligenza artificiale" not in page
assert "non documenta" not in page
assert page.count('<table class="data-table">') == 2
assert page.count('<h2 id="sezione-') == 15
assert page.count('<details class="dossier-contents">') == 1
assert D["image"] in page and D["inlineImage"] in page and D["socialImage"] in page
assert D["title"] in page
assert "PUTTANA IVA!" not in page

for rejected in (
    "Agnelli", "Lamborghini", "Rovagnati", "Il vestito non serve",
    "Il passato non si resetta", "L’IVA respira", "Figura 2",
):
    assert rejected not in page, rejected

sources_html = re.search(r'<ul class="sources-list">(.*?)</ul>', page, re.S).group(1)
assert sources_html.count("<li>") >= 20
assert "<a " not in sources_html
assert "http://" not in sources_html and "https://" not in sources_html

plain = normalize(html.unescape(re.sub(r"<[^>]+>", " ", page)))
document = Document(ROOT / D["sourceDocx"])
missing = []
for paragraph in document.paragraphs[6:]:
    value = normalize(paragraph.text)
    if value and value != "FONTI" and paragraph.style.name != "Caption" and value not in plain:
        missing.append(value)
assert not missing, "Missing DOCX paragraphs: " + " | ".join(missing[:5])
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            assert normalize(cell.text) in plain, f"Missing table text: {cell.text}"

for expanded in (
    "Imposta sul valore aggiunto (IVA)", "Imposta generale sull’entrata (IGE)",
    "Comunità economica europea (CEE)", "Istituto nazionale di statistica (ISTAT)",
    "Istituto per la Ricostruzione Industriale (IRI)",
    "Società Meridionale di Elettricità (SME)",
):
    assert expanded in plain, expanded

for filename in ("index.html", "indagini.html", "archivio-economia.html", "sitemap-articles.xml", "sitemap-news.xml", "feed.xml"):
    content = (ROOT / filename).read_text(encoding="utf-8")
    assert D["slug"] in content, f"Missing placement in {filename}"
    assert content.count(D["slug"]) >= 1

home = (ROOT / "index.html").read_text(encoding="utf-8")
lead = re.search(r'<article class="lead-story".*?</article>', home, re.S).group(0)
assert D["slug"] in lead and 'fetchpriority="high"' in lead

indagini = (ROOT / "indagini.html").read_text(encoding="utf-8")
economy = re.search(r'<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-economia.html">Economia</a></h2></div><div class="archive-grid">(.*?)</div><a class="category-archive-link" href="archivio-economia.html">', indagini, re.S).group(1)
assert economy.count('<article class="archive-card"') == 3
assert economy.index(D["slug"]) < economy.index("article-economia-energia-eni-price-cap-carburanti.html")

archive = (ROOT / "archivio-economia.html").read_text(encoding="utf-8")
archive_grid = re.search(r'<div class="archive-grid">(.*?)</div><footer', archive, re.S).group(1)
assert archive_grid.index(D["slug"]) < archive_grid.index("article-economia-energia-eni-price-cap-carburanti.html")
assert archive_grid.count(f'data-dossier="{D["slug"]}"') == 1

ids = set(re.findall(r'\bid="([^"]+)"', page))
for value in re.findall(r'(?:href|src)="([^"]+)"', page):
    if value.startswith("#"):
        assert value[1:] in ids, f"Broken fragment: {value}"
    elif value and not value.startswith(("http://", "https://", "mailto:", "/")):
        assert (ROOT / value.split("?", 1)[0]).exists(), f"Broken local reference: {value}"

schema = json.loads(re.search(r'<script type="application/ld\+json" data-seo-schema>(.*?)</script>', page, re.S).group(1))
article = next(item for item in schema["@graph"] if item.get("@type") == "NewsArticle")
assert article["headline"] == D["title"]
assert article["datePublished"] == D["datePublished"]
assert article["dateModified"] == D["dateModified"]
assert article["articleSection"] == D["macroCategory"]
assert article["image"] == ["https://unosguardosulluomo.github.io/" + D["socialImage"]]

for sitemap in ("sitemap.xml", "sitemap-google.xml", "sitemap-index.xml", "sitemap-pages.xml", "sitemap-articles.xml", "sitemap-news.xml"):
    tree = ET.parse(ROOT / sitemap)
articles_xml = (ROOT / "sitemap-articles.xml").read_text(encoding="utf-8")
assert f"<loc>{url}</loc><lastmod>{D['datePublished']}</lastmod>" in articles_xml
assert "Sitemap:" in (ROOT / "robots.txt").read_text(encoding="utf-8")

print("IVA 1973 dossier checks passed")
