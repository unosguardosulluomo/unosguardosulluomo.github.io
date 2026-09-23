"""Prepare the definitive Giorgetti dossier while preserving its voice and layout."""

from pathlib import Path
import shutil
import sys

from docx import Document


REPLACEMENTS = {
    "Il 22 settembre 2026 l’Istat ha confermato il deficit italiano del 2025 al 3,1% del PIL. Il numero, preso da solo, non descrive un crollo. Al contrario: il disavanzo è sceso rispetto agli anni peggiori, il saldo primario è tornato positivo e il percorso di risanamento esiste. Ma proprio per questo il caso è più sottile. Il problema non è che l’Italia sia precipitata fuori strada. Il problema è che, pochi giorni prima della certificazione, il governo aveva lasciato intravedere la possibilità di un’uscita anticipata dalla procedura europea per disavanzo eccessivo.":
    "Il 22 settembre 2026 l’Istituto nazionale di statistica (Istat), nella revisione dei Conti economici nazionali, ha confermato il deficit italiano del 2025 al 3,1% del prodotto interno lordo (PIL). Il numero, preso da solo, non descrive un crollo. Al contrario: il disavanzo è sceso rispetto agli anni peggiori, il saldo primario è tornato positivo e il percorso di risanamento esiste. Ma proprio per questo il caso è più sottile. Il problema non è che l’Italia sia precipitata fuori strada. Il problema è che, pochi giorni prima della revisione, il governo aveva lasciato intravedere la possibilità di un’uscita anticipata dalla procedura europea per disavanzo eccessivo.",
    "Giancarlo Giorgetti non aveva promesso formalmente il risultato. Aveva scelto una formula più prudente: si era detto «speranzoso», rimettendo il verdetto all’Istat. Sul piano letterale è una frase cauta. Sul piano politico, però, una speranza dichiarata dal ministro dell’Economia a quattro giorni da una certificazione che vale miliardi non è una frase neutra. È una aspettativa pubblica.":
    "Giancarlo Giorgetti non aveva promesso formalmente il risultato. Aveva scelto una formula più prudente: si era detto «speranzoso», rimettendo il primo verdetto all’Istat. La revisione statistica non avrebbe chiuso da sola la procedura: il passaggio formale sarebbe spettato alle istituzioni europee. Sul piano letterale, quindi, la sua era una frase cauta. Sul piano politico, però, una speranza dichiarata dal ministro dell’Economia a quattro giorni da una revisione dei conti che vale miliardi non è una frase neutra. È un’aspettativa pubblica.",
    "Se quelle maggiori entrate non fossero arrivate, a parità di tutto il resto, il deficit sarebbe stato nell’ordine del 3,6% del PIL. Quindi una parte importante del miglioramento effettivo non deriva soltanto da tagli o disciplina, ma anche dal fatto che le entrate hanno tenuto più del previsto.":
    "Se si prende il solo scostamento delle entrate e si lascia fermo tutto il resto — un esercizio aritmetico, non una ricostruzione economica — il deficit sarebbe stato nell’ordine del 3,6% del PIL. Il calcolo non dice che quello sarebbe stato lo scenario reale: alcune entrate erano legate a poste contabili accompagnate da spese corrispondenti. Dice però quanto le entrate superiori alle previsioni abbiano contribuito a sostenere il saldo.",
    "Il punto non è ridurre tutto all’IVA sulla benzina. Sarebbe troppo piccolo. La vera questione è macroeconomica.":
    "Il punto non è ridurre tutto all’imposta sul valore aggiunto (IVA) sulla benzina. Sarebbe troppo piccolo. La vera questione è macroeconomica.",
    "Durante questa indagine è emersa una tentazione naturale: mettere accanto i miliardi del decreto sul bollo auto e i pochi miliardi che separavano l’Italia dalla soglia utile, e concludere che bastava spostare quei soldi.":
    "Durante questa indagine è emersa una tentazione naturale: mettere accanto i 2,2935 miliardi destinati dal decreto-legge 17 settembre 2026, n. 162, alla compensazione dell’esenzione del bollo auto per il 2027 e i pochi miliardi che separavano l’Italia dalla soglia utile, e concludere che bastava spostare quei soldi.",
}


SOURCES = [
    "Istituto nazionale di statistica (Istat) — Conti economici nazionali, anni 2010-2025, 22 settembre 2026.",
    "Istat — Notifica dell’indebitamento netto e del debito delle Amministrazioni pubbliche secondo il Trattato di Maastricht, anni 2022-2025, 22 aprile 2026.",
    "Istat — PIL e indebitamento delle Amministrazioni pubbliche, anni 2023-2025, 2 marzo 2026.",
    "Ufficio parlamentare di bilancio — Audizione sul Documento di finanza pubblica 2026, 28 aprile 2026.",
    "Ufficio parlamentare di bilancio — Rapporto sulla politica di bilancio 2026, 10 giugno 2026.",
    "Consiglio dell’Unione europea (UE) — Patto di stabilità e crescita: raccomandazioni ai paesi sottoposti a procedura per disavanzo eccessivo, 21 gennaio 2025.",
    "Commissione europea — European Economic Forecast, Spring 2026, 21 maggio 2026.",
    "Banca d’Italia — Relazione annuale sul 2025, 29 maggio 2026.",
    "Gazzetta Ufficiale della Repubblica Italiana — Decreto-legge 17 settembre 2026, n. 162, Serie Generale n. 216, 17 settembre 2026.",
    "Agenzia Nazionale Stampa Associata (ANSA) — Giorgetti, ‘sono speranzoso sull’uscita dalla procedura UE per disavanzo’, 18 settembre 2026.",
    "ANSA — Giorgetti, ‘rammarico, l’Italia non esce anticipatamente dalla procedura di deficit’, 22 settembre 2026.",
]


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def revise(source: Path, repo_output: Path, delivery_output: Path) -> None:
    doc = Document(source)
    found = set()
    for paragraph in doc.paragraphs:
        replacement = REPLACEMENTS.get(paragraph.text)
        if replacement is not None:
            found.add(paragraph.text)
            paragraph.text = replacement

    missing = set(REPLACEMENTS) - found
    if missing:
        raise RuntimeError(f"Missing expected source paragraphs: {len(missing)}")

    paragraphs = list(doc.paragraphs)
    source_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Fonti principali")
    for paragraph in paragraphs[source_index + 1:]:
        remove_paragraph(paragraph)
    for item in SOURCES:
        doc.add_paragraph(item, style="Normal")

    doc.core_properties.title = "Giorgetti e il pallottoliere"
    doc.core_properties.subject = "Dossier sui conti pubblici italiani e sulla comunicazione politica del deficit 2025"
    doc.core_properties.author = "Redazione Uno Sguardo sull’Uomo"
    doc.core_properties.keywords = "Giorgetti, deficit, conti pubblici, Istat, debito pubblico, procedura per disavanzo eccessivo"

    repo_output.parent.mkdir(parents=True, exist_ok=True)
    delivery_output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(repo_output)
    shutil.copy2(repo_output, delivery_output)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Usage: revise-giorgetti-docx.py SOURCE REPO_OUTPUT DELIVERY_OUTPUT")
    revise(*(Path(value) for value in sys.argv[1:]))
