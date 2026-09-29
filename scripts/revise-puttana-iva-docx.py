"""Create the definitive Puttana IVA dossier from the narrative draft."""

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


def args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--shop-image", required=True)
    parser.add_argument("--bar-image", required=True)
    return parser.parse_args()


def split_sections(document):
    sections = {}
    current = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name == "Heading 1":
            if text == "FONTI":
                current = None
            else:
                current = text
                sections[current] = []
        elif current and text:
            sections[current].append(text)
    return sections


def by_prefix(sections, prefix):
    return next(sections[key] for key in sections if key.startswith(prefix))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "D9D9D9")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[index]
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F3F6F9")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(9)
    document.add_paragraph()
    return table


def add_picture(document, path, caption, width=Inches(6.6)):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=width)
    cap = document.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True


def add_body(document, text, bold=False):
    p = document.add_paragraph()
    p.paragraph_format.keep_together = False
    run = p.add_run(text)
    run.bold = bold
    return p


def add_section(document, title, paragraphs):
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        add_body(document, paragraph)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.13

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(13)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(14)

    heading = styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(70, 70, 70)

    bullet = styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(9.5)
    bullet.paragraph_format.space_after = Pt(3)


def main():
    options = args()
    source = Document(options.input)
    old = split_sections(source)
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    core = document.core_properties
    core.title = "1973: come l’IVA cambiò la vita degli italiani"
    core.subject = "Storia dell’IVA, trasformazione dei prezzi e vulnerabilità informativa dei consumatori"
    core.author = "UNO SGUARDO SULL’UOMO"
    core.keywords = "IVA, IGE, prezzi, inflazione, consumatori, euro, 1973, economia"
    core.comments = "Dossier editoriale definitivo"

    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("1973: come l’IVA cambiò la vita degli italiani")
    p = document.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Dal 1940 al 2026: come una riforma necessaria cambiò il modo di leggere i prezzi")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UNO SGUARDO SULL’UOMO")
    r.bold = True
    r.font.size = Pt(11)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Dossier definitivo — Economia")

    add_picture(
        document,
        Path(options.shop_image),
        "Immagine prodotta da Uno Sguardo sull’Uomo.",
    )
    standfirst = (
        "Nel 1973 l’Imposta sul valore aggiunto (IVA) entrò in una società con una scolarizzazione molto più bassa, "
        "conoscenze concentrate nei singoli mestieri e un’informazione difficile da verificare. Nel 2002 il passaggio "
        "all’euro produsse un disorientamento diverso ma simile: per qualche tempo il consumatore perse la capacità di "
        "giudicare un prezzo a colpo d’occhio. Questa è la storia di ciò che accade quando le regole cambiano più in fretta "
        "degli strumenti disponibili per comprenderle."
    )
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(standfirst)
    r.bold = True
    r.font.size = Pt(11)
    p.add_run().add_break(WD_BREAK.PAGE)

    add_section(document, "1. Ricordi il 2002", [
        "Chi oggi ha cinquantacinque o sessant’anni difficilmente ricorda che cosa accadde il 1° gennaio 1973. Era bambino, oppure non era ancora nato. Ma il 2002 lo ricorda.",
        "Ricorda i cartellini riscritti e il riflesso di trasformare mentalmente gli euro in lire. Il prezzo era davanti agli occhi, ma per qualche secondo non diceva più niente.",
        "Durante questa indagine è tornato un ricordo preciso: un paio di scarpe da 50.000 lire che, dopo il cambio, compare a 50 euro. Il cambio ufficiale avrebbe trasformato 50.000 lire in circa 25,82 euro. È una memoria individuale, non una misura dell’inflazione. Mostra però con chiarezza quanto fosse difficile, nei primi mesi, stabilire se un nuovo prezzo fosse normale, alto o sproporzionato.",
        "Quando il riferimento abituale scompare, il consumatore perde temporaneamente una parte della propria capacità di giudizio. Nel 1973 era accaduto qualcosa di diverso, ma con un effetto simile: la moneta era rimasta la stessa, mentre era cambiato il sistema fiscale incorporato nel prezzo.",
    ])

    sec2 = []
    for prefix in ("2.", "3.", "4."):
        sec2.extend(by_prefix(old, prefix))
    sec2 = [p for p in sec2 if not p.startswith(("IGE: una tassa", "Con l’IGE l’imposta"))]
    sec2[0] = sec2[0].replace("l’Imposta Generale sull’Entrata, l’IGE", "l’Imposta generale sull’entrata (IGE)")
    add_section(document, "2. Prima dell’IVA c’era l’IGE", sec2)

    sec3 = []
    for prefix in ("5.", "6."):
        sec3.extend(by_prefix(old, prefix))
    sec3 = [p for p in sec3 if p != "La CEE uniformò la macchina, non il peso dell’imposta."]
    sec3[0] = sec3[0].replace("Comunità Economica Europea", "Comunità economica europea (CEE)")
    add_section(document, "3. Perché l’Europa scelse l’IVA", sec3)

    add_section(document, "4. L’Italia del 1971 e la conoscenza divisa per mestiere", [
        "Due anni prima dell’arrivo dell’IVA, il censimento del 1971 fotografa un’Italia molto diversa da quella di oggi. Fra le persone di almeno sei anni, il 5,2 per cento è analfabeta e il 27,1 per cento sa leggere e scrivere ma non possiede alcun titolo di studio. Il 44,3 per cento ha la licenza elementare. Il 14,7 per cento ha concluso la scuola media, il 6,9 per cento possiede un diploma e l’1,8 per cento una laurea.",
        "I dati dell’Istituto nazionale di statistica (ISTAT) non misurano l’intelligenza delle persone e non descrivono la loro capacità di lavorare. Misurano la distanza fra il linguaggio della riforma e gli strumenti scolastici con cui una parte enorme della popolazione doveva affrontarlo. Più di tre persone su quattro non erano andate oltre la scuola elementare oppure non possedevano alcun titolo.",
        "L’Italia del 1973 non era priva di competenze. Le competenze, però, rimanevano più spesso confinate al mestiere. Il contadino conosceva la terra, gli animali e le stagioni. L’artigiano conosceva i materiali. Il geometra conosceva il cantiere. Il farmacista conosceva i medicinali. Un sapere poteva essere profondo senza comprendere la contabilità di un altro settore.",
        "L’imponibile, la detrazione, l’IVA a credito e l’IVA a debito erano concetti nuovi anche per molte persone istruite che non avevano mai tenuto un registro fiscale. Per chi possedeva soltanto l’istruzione elementare, o non sapeva leggere e scrivere, la difficoltà era maggiore. La riforma introduceva parole, calcoli e procedure che la scuola, il mestiere e la vita quotidiana non avevano avuto ragione di insegnare.",
        "Anche l’informazione funzionava in modo diverso. Nel 1973 gli abbonamenti televisivi privati erano circa 11,7 milioni, pari a 72,5 ogni cento famiglie. Il televisore era già un mezzo di massa, ma offriva due reti pubbliche, orari stabiliti e un palinsesto scelto dall’alto. Il cittadino poteva ricevere un’informazione; non poteva cercarla, confrontarla con decine di fonti o mettere immediatamente in comune la propria esperienza con quella di altri consumatori.",
        "Se il telegiornale e il quotidiano non spiegavano la nuova imposta in modo comprensibile, restavano il ragioniere, il commerciante, il vicino di casa e l’esperienza diretta. Milioni di persone stavano per incontrare il nuovo sistema sul lavoro, nella fattura o sul cartellino, senza disporre di un luogo nel quale verificare subito ciò che veniva loro detto.",
    ])

    sec5 = []
    for prefix in ("9.", "10.", "11.", "12."):
        sec5.extend(by_prefix(old, prefix))
    sec5 = [p for p in sec5 if p != "Il consumatore vede il nuovo cartellino. E basta."]
    add_section(document, "5. Le settimane in cui tutti dovettero rifare i prezzi", sec5)

    add_section(document, "6. Il caffè porta la riforma dentro il bar", [
        "Per vedere la distanza fra la riforma fiscale e la vita quotidiana basta entrare in un bar.",
        "Una cronaca laziale dei primi giorni di gennaio 1973 registra il caffè che passa da circa 70 lire a nuovi listini fra 90 e 110 lire. Per il cliente non esistono direttive europee, imponibili o meccanismi di detrazione. Esiste una tazzina che ieri costava 70 lire e oggi ne costa 90, 100 o 110.",
        "Il cliente può chiedere una spiegazione al barista e confrontare il prezzo con gli esercizi vicini. Non può verificare in pochi minuti che cosa stia accadendo in altre città, leggere la norma, consultare spiegazioni concorrenti o scoprire che migliaia di persone stanno segnalando la stessa variazione.",
        "La riforma diventa reale quando un gesto automatico cambia prezzo e chi compra non dispone ancora di un riferimento per distinguere il nuovo carico fiscale, l’aumento dei costi e la decisione commerciale.",
    ])
    add_picture(
        document,
        Path(options.bar_image),
        "Immagine prodotta da Uno Sguardo sull’Uomo.",
    )

    add_body(document, "Nella finestra ottobre 1972-gennaio 1973 i listini all’ingrosso di diversi alimentari si muovono mentre gli operatori preparano il passaggio dall’IGE all’IVA. La coincidenza temporale documenta il riprezzamento, ma non consente di attribuire automaticamente alla nuova imposta ogni aumento osservato.")
    add_table(document, ["Prodotto", "Variazione", "Livello e periodo"], [
        ("Pomodori pelati", "+40%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Olio d’oliva", "+28-30%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Concentrato di pomodoro", "+25%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Liquori", "+12-15%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Caffè", "+10-12%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Vino", "+10-12%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Carne suina fresca", "+5-7%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Biscotti", "+3-5%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Tonno e sardine", "+3-5%", "Ingrosso, ottobre 1972-gennaio 1973"),
        ("Salumi e insaccati", "+3-5%", "Ingrosso, ottobre 1972-gennaio 1973"),
    ], [Inches(2.0), Inches(1.25), Inches(3.35)])
    add_body(document, "Un secondo gruppo di dati riguarda livelli e periodi differenti. Non deve essere confrontato come se descrivesse un’unica serie: serve a mostrare che la pressione sui prezzi attraversava produzione, ingrosso e consumo in modi diversi.")
    add_table(document, ["Bene", "Variazione", "Livello"], [
        ("Carne bovina", "+2,8%", "Prezzo al consumo"),
        ("Vitello", "+9,4%", "Prezzo alla produzione"),
        ("Maiale", "+2,0%", "Prezzo al consumo"),
        ("Pollo", "+12,8%", "Prezzo alla produzione"),
        ("Pesce comune", "+10,6%", "Prezzo al consumo"),
        ("Latte, formaggi e uova", "+6,3%", "Categoria aggregata al consumo"),
        ("Caffè", "+5,2%", "Consumo; +10-12% all’ingrosso ottobre-gennaio"),
        ("Conserve di pomodoro", "Pelati +40%; concentrato +25%", "Ingrosso ottobre-gennaio"),
        ("Uova", "+3,1%", "Prezzo al consumo"),
        ("Pane", "Circa 0%", "Media 1972-1973"),
        ("Olio d’oliva", "+21,7%", "Consumo; +28-30% all’ingrosso ottobre-gennaio"),
    ], [Inches(2.0), Inches(1.9), Inches(2.7)])
    add_body(document, "Farina e legumi restano fuori perché nella documentazione raccolta non è stata trovata una serie omogenea abbastanza solida da trattarli insieme agli altri prodotti.")

    sec7 = []
    for prefix in ("15.", "16.", "17."):
        sec7.extend(by_prefix(old, prefix))
    sec7 = [p for p in sec7 if not p.startswith(("L’IVA non costava", "Quando il prezzo si riscrive"))]
    sec7 = [p.replace("Cirio è nell’orbita IRI-SME.", "Cirio è nell’orbita dell’Istituto per la Ricostruzione Industriale (IRI) e della Società Meridionale di Elettricità (SME).") for p in sec7]
    sec7.append("La documentazione disponibile mostra quindi una filiera in movimento durante il cambio d’imposta. Dimostra l’aumento pagato dal consumatore, ma non permette di assegnare con precisione ogni parte di quel rincaro all’IVA, alle materie prime, agli imballaggi o ai margini commerciali.")
    add_section(document, "7. Il caso dei pelati e la filiera che il consumatore non vede", sec7)

    sec8 = by_prefix(old, "18.")
    sec8 = [p for p in sec8 if not p.startswith("La famiglia non mangiava")]
    sec8.insert(4, "Per una famiglia, però, il dato decisivo non è la media astratta: è il prezzo dei beni acquistati ogni settimana. Pane, pasta, carne, olio, uova e pomodori pesano sulla percezione molto più di una voce stabile o in diminuzione comprata raramente.")
    add_section(document, "8. Perché la media dell’inflazione non descrive ogni famiglia", sec8)

    add_section(document, "9. Campagna e città non incontrano gli stessi prezzi", [
        "L’Italia del 1973 è già industriale e urbanizzata, ma il rapporto con l’agricoltura e con l’autoproduzione è molto più esteso di oggi. Il censimento agricolo del 1970 rileva circa 3,6 milioni di aziende. I piccoli orti e gli allevamenti familiari destinati all’autoconsumo non rientrano neppure nel campo di osservazione del censimento.",
        "Molte famiglie mantengono un orto, alcune galline, un piccolo appezzamento o parenti in paese. Uova, polli, ortaggi, conserve, vino e olio possono arrivare in casa senza attraversare ogni volta l’intera filiera commerciale.",
        "Questo non rende immuni dall’inflazione. Aumentano anche mangimi, carburanti, attrezzi e beni che la famiglia non può produrre. Ma chi raccoglie le uova dal proprio pollaio non incontra ogni settimana il loro nuovo prezzo sul banco del negozio.",
        "Nelle città industriali la situazione è diversa. La famiglia operaia vive in appartamento, riceve un salario monetario e compra quasi tutto. Il pane si compra, l’olio si compra, la carne si compra. Ogni aumento arriva direttamente alla cassa.",
        "Nel 1973 alimentari e bevande assorbono circa un terzo della spesa familiare; nelle famiglie di lavoratori dipendenti il peso è ancora maggiore. Per questo una variazione a doppia cifra sui beni essenziali può modificare la vita anche quando l’indice generale cresce meno.",
        "La rinuncia prende forme concrete: meno carne, una qualità inferiore, un acquisto rimandato, un’abitudine cancellata. Anche i salari crescono, ma la famiglia vive nel tempo. Se il prezzo sale a gennaio e il recupero salariale arriva mesi dopo, la spesa cambia prima della busta paga.",
    ])

    add_section(document, "10. Nel 2002 la stessa perdita di orientamento cambia moneta", [
        "Nel 2002 l’italiano ha un punto fermo: un euro vale 1.936,27 lire. Può fare la conversione, ma deve costruire una nuova memoria dei prezzi. Per qualche tempo il numero sul cartellino non produce più una valutazione immediata.",
        "La televisione è ormai universale e l’informazione è molto più ampia rispetto al 1973. Internet e i telefoni cellulari esistono, ma non sono ancora strumenti permanenti nelle mani di tutti. Non esistono social network di massa e confrontare prezzi online non è un gesto quotidiano. Anche nel 2002 la comunicazione resta in gran parte verticale.",
        "La Banca d’Italia osserva che, dovendo produrre nuovi listini, le imprese possono utilizzare il cambio di moneta per anticipare o rinviare adeguamenti. Uno studio successivo su 2.500 ristoranti e trattorie rileva che, al momento dell’introduzione dell’euro, la quota di imprese che rivede i listini sale al 75 per cento.",
        "Le statistiche ufficiali non mostrano un raddoppio generalizzato dei prezzi: per l’area euro, l’effetto aggregato del cambio del contante nella prima metà del 2002 viene stimato fra zero e 0,2 punti percentuali. Registrano però aumenti in alcuni beni e servizi e una distanza eccezionale fra inflazione misurata e inflazione percepita.",
        "La somiglianza con il 1973 non consiste quindi nell’attribuire ogni rincaro all’IVA o all’euro. Consiste nel fatto che migliaia di prezzi vengono riscritti mentre il consumatore ha temporaneamente perduto la propria unità di confronto. Chi prepara il listino conosce costi, margini e vecchio prezzo. Chi compra vede il numero finale.",
    ])

    sec11 = []
    for prefix in ("24.", "25."):
        sec11.extend(by_prefix(old, prefix))
    sec11 = [
        p for p in sec11
        if not p.startswith(("L’IVA è una tassa che respira", "È anche per questo", "Figura 2"))
    ]
    add_section(document, "11. Quando l’IVA diventa una parte stabile del bilancio pubblico", sec11)

    add_section(document, "12. Quanto costa davvero ridurre l’IVA", [
        "“Abbassiamo l’IVA” è una proposta immediatamente comprensibile: il cittadino immagina un prezzo più basso. Per valutarla, però, bisogna distinguere una riduzione di cinque punti dell’aliquota ordinaria da una riduzione del cinque per cento del gettito complessivo. Non sono la stessa operazione.",
        "Nel bilancio dello Stato per il 2025 la previsione di cassa dell’IVA è di circa 192,7 miliardi di euro. Questa cifra comprende entrate prodotte da basi imponibili e aliquote differenti. Non permette, da sola, di calcolare il costo di un passaggio dell’aliquota ordinaria dal 22 al 17 per cento.",
        "Per stimare quel costo servono almeno la base imponibile effettivamente assoggettata all’aliquota ordinaria, la quota di riduzione trasferita ai prezzi, la reazione dei consumi e gli effetti sulle altre entrate. Il dato di 192,7 miliardi mostra la dimensione della leva fiscale; non è il preventivo di una proposta.",
        "Il carburante rende visibile il meccanismo perché nel prezzo convivono componente industriale, accisa e IVA. L’accisa entra nella base imponibile sulla quale viene calcolata l’IVA. Se aumenta il prezzo industriale o cambia l’accisa, cambia anche l’IVA in valore assoluto.",
        "Ridurre l’imposta può diminuire il prezzo solo nella misura in cui il taglio viene trasferito al consumatore. Contemporaneamente riduce un’entrata pubblica. La scelta completa deve quindi specificare durata, prodotti interessati, costo previsto, comportamento atteso dei prezzi e copertura.",
    ])

    add_section(document, "13. Lo Stato eredita impegni e deve scegliere come finanziarli", [
        "Uno Stato non chiude i conti il 31 dicembre per ripartire da zero il 1° gennaio. Ogni generazione eredita pensioni, sanità, scuola, infrastrutture, enti, trasferimenti, debiti, interessi, diritti, contratti e progetti.",
        "Alcune decisioni hanno costruito il Paese, altre hanno prodotto inefficienze e altre ancora sono state errori. Cambiare governo non cancella gli impegni finanziari già assunti né il costo delle decisioni accumulate.",
        "Ridurre una grande imposta significa modificare un equilibrio esistente: meno spesa, maggiori entrate altrove, più debito oppure una combinazione di queste possibilità. La discussione diventa concreta quando indica quali servizi, trasferimenti, investimenti o altre entrate devono cambiare.",
        "Il bilancio pubblico è il punto nel quale decisioni prese in tempi diversi vengono pagate insieme. Discutere l’IVA significa quindi discutere anche ciò che quella entrata finanzia e chi sopporta il costo della sua eventuale riduzione.",
    ])

    sec14 = by_prefix(old, "29.")
    add_section(document, "14. I limiti europei di oggi", sec14)

    add_section(document, "15. Cinquantatré anni dopo", [
        "L’IVA non nasce dal nulla nel 1973. Arriva dopo trentatré anni di IGE e dopo la nascita di un mercato comune che ha bisogno di conoscere quanta imposta viaggia dentro ogni prodotto.",
        "Per l’Europa è una modernizzazione necessaria. Per lo Stato italiano è una rivoluzione contabile. Per milioni di persone è un cambio di linguaggio che arriva in poche settimane, dentro una società con livelli di istruzione molto diseguali e pochi strumenti per verificare rapidamente ciò che accade ai prezzi.",
        "Il consumatore del 1973 non è ingenuo davanti alla vita. È inesperto davanti a una macchina fiscale che fino al giorno prima non aveva avuto motivo di conoscere. Chi vive ancora in parte di autoconsumo può sottrarre alcuni beni al mercato. La famiglia urbana salariata incontra invece ogni aumento alla cassa.",
        "Nel 2002 il sistema informativo è più ricco, ma non ancora continuo e partecipato come oggi. Cambia la moneta e il consumatore perde nuovamente, per qualche tempo, la capacità di riconoscere il valore abituale del prezzo.",
        "In entrambe le transizioni, chi costruisce i listini possiede informazioni che chi compra non ha. Questo non dimostra che ogni aumento dipenda dall’IVA o dall’euro. Spiega perché una modifica simultanea di migliaia di prezzi possa produrre disorientamento, sospetto e rinunce anche quando la media statistica racconta un movimento più contenuto.",
        "Dopo cinquantatré anni l’IVA è diventata una delle grandi entrate pubbliche italiane. La parola è familiare; il meccanismo e le conseguenze delle sue variazioni molto meno. Comprenderli richiede ancora di separare l’imposta dai costi, i costi dai margini, la percezione dalla media e la promessa politica dal conto necessario a sostenerla.",
    ])

    document.add_heading("FONTI", level=1)
    sources = [
        "Regno d’Italia, Regio decreto-legge 9 gennaio 1940, n. 2, “Istituzione di una imposta generale sull’entrata”; entrata in vigore 8 febbraio 1940.",
        "Regno d’Italia, Legge 19 giugno 1940, n. 762, conversione del Regio decreto-legge 9 gennaio 1940, n. 2.",
        "Comunità economica europea, Prima direttiva del Consiglio 67/227/CEE, 11 aprile 1967, armonizzazione delle legislazioni degli Stati membri relative alle imposte sulla cifra d’affari.",
        "Comunità economica europea, Seconda direttiva del Consiglio 67/228/CEE, 11 aprile 1967, struttura e modalità di applicazione del sistema comune di imposta sul valore aggiunto.",
        "Repubblica Italiana, Legge 9 ottobre 1971, n. 825, delega legislativa al Governo per la riforma tributaria.",
        "Repubblica Italiana, Decreto del Presidente della Repubblica 26 ottobre 1972, n. 633, “Istituzione e disciplina dell’imposta sul valore aggiunto”, testo originario e disposizioni transitorie.",
        "Camera dei Deputati, VI Legislatura, Discussioni, seduta del 4 aprile 1973, interventi sugli effetti dell’IVA sui prezzi e stima degli uffici della programmazione.",
        "Banca d’Italia, “Relazione annuale sul 1973” e Appendice, 31 maggio 1974.",
        "Istituto nazionale di statistica, Censimento generale della popolazione 1971, popolazione residente di almeno sei anni per grado di istruzione.",
        "Istituto nazionale di statistica, “L’Italia in 150 anni. Sommario di statistiche storiche 1861-2010”, capitoli 7, 8 e 13, 2011.",
        "Istituto nazionale di statistica, Censimenti storici dell’agricoltura, risultati dei censimenti 1961 e 1970.",
        "Istituto nazionale di statistica, “I cambiamenti dell’agricoltura”, Storie di dati, 12 giugno 2026.",
        "Istituto nazionale di statistica, “I consumi cambiano insieme al Paese”, Storie di dati, 2026.",
        "Archivio storico de l’Unità, cronaca e rilevazioni sugli aumenti all’ingrosso nella fase di introduzione dell’IVA, 1° aprile 1973.",
        "Lotta Continua, cronache sui listini dei bar e sul prezzo del caffè, gennaio 1973.",
        "Camera dei Deputati e Senato della Repubblica, atti parlamentari 1972-1973 sulla transizione IVA, i prezzi e l’industria conserviera.",
        "Consiglio per la ricerca in agricoltura e l’analisi dell’economia agraria e fonti storiche agricole, dati sulla campagna del pomodoro 1972-1973.",
        "Dipartimento delle Finanze, documentazione storica sull’evoluzione dell’aliquota ordinaria IVA.",
        "Repubblica Italiana, normativa sull’aumento dell’aliquota ordinaria IVA al 22 per cento dal 1° ottobre 2013.",
        "Ministero dell’Economia e delle Finanze, Disegno di legge di bilancio 2025, relazione, articolato e quadri generali; previsione di cassa IVA pari a 192,717 miliardi di euro.",
        "Unione europea, Direttiva 2006/112/CE relativa al sistema comune d’imposta sul valore aggiunto, testo consolidato.",
        "Unione europea, Direttiva (UE) 2022/542 del Consiglio, riforma delle aliquote IVA, 5 aprile 2022.",
        "Banca d’Italia, “L’impatto del cambio della moneta sui prezzi: prime valutazioni”, Bollettino economico n. 38, marzo 2002.",
        "Banca d’Italia, Fabiani, Gattulli e Sabbatini, “L’introduzione dell’euro e le politiche di prezzo: analisi di un campione di dati individuali”, Temi di discussione n. 541, dicembre 2004.",
        "Banca d’Italia, Del Giovane e Sabbatini, “L’introduzione dell’euro e la divergenza tra inflazione rilevata e percepita”, Temi di discussione n. 532, dicembre 2004.",
        "Banca centrale europea, “Rapporto annuale 2002”, sezione sul cambio del contante e sull’inflazione percepita, 2003.",
    ]
    for source_text in sources:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(source_text)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.add_run("UNO SGUARDO SULL’UOMO — Dossier Economia")

    output = Path(options.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
