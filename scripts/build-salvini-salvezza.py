"""Build the approved dossier on Salvini, Zaia and the Lega congress."""

import email.utils
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/salvini-salvezza-lega.json").read_text(encoding="utf-8"))
SITE = "https://unosguardosulluomo.github.io/"
URL = SITE + D["slug"]


def esc(value):
    return html.escape(str(value), quote=True)


def rich_paragraph(paragraph):
    fragments = []
    for run in paragraph.runs:
        value = esc(run.text).replace("\n", "<br>")
        if not value:
            continue
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments) or esc(paragraph.text)


document = Document(ROOT / D["sourceDocx"])
body = []
contents = []
sources = []
source_intro = (
    "Le fonti sono riportate per consentire il controllo dei dati e delle ricostruzioni. "
    "Nel corpo del dossier non sono inseriti richiami numerici: i documenti sostengono "
    "i fatti verificabili, non sostituiscono l'analisi editoriale."
)
criterion = (
    "Il dossier distingue il voto dei delegati dal consenso elettorale, i poteri previsti "
    "dallo statuto dalle loro conseguenze politiche e le ricostruzioni giornalistiche dai "
    "fatti documentati. Le intenzioni personali non accessibili restano interpretazioni e "
    "i sondaggi non sono presentati come risultati elettorali."
)
started = False
in_sources = False
heading_number = 0
lead = ""
list_open = False


def close_list():
    global list_open
    if list_open:
        body.append("</ul>")
        list_open = False


for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if not text:
        continue
    style = paragraph.style.name
    if not started:
        if style == "Dossier Body" and text.startswith("Matteo Salvini ha vinto il congresso della Lega"):
            started = True
            lead = rich_paragraph(paragraph)
        continue
    if style == "Heading 1" and re.sub(r"^\d+\.\s*", "", text) == "Fonti documentali":
        close_list()
        in_sources = True
        continue
    if in_sources:
        sources.append(text.removeprefix("•").strip())
        continue
    if style == "Caption":
        continue
    if not lead:
        lead = rich_paragraph(paragraph)
        continue
    if style == "Heading 1":
        close_list()
        display_heading = re.sub(r"^\d+\.\s*", "", text)
        if display_heading == "Zaia non è il capo del ventiquattro per cento":
            body.append(
                '<figure class="article-inline-image">'
                f'<img src="{D["inlineImage"]}" width="{D["inlineImageWidth"]}" height="{D["inlineImageHeight"]}" '
                f'alt="{esc(D["inlineImageAlt"])}" loading="lazy">'
                f'<figcaption>{esc(D["imageCaption"])}</figcaption></figure>'
            )
        heading_number += 1
        anchor = f"sezione-{heading_number}"
        contents.append(f'<li><a href="#{anchor}">{esc(display_heading)}</a></li>')
        body.append(f'<h2 id="{anchor}">{esc(display_heading)}</h2>')
    elif style == "Heading 2":
        close_list()
        body.append(f'<h3>{esc(text)}</h3>')
    elif style.startswith("List"):
        if not list_open:
            body.append('<ul class="article-list">')
            list_open = True
        body.append(f'<li>{rich_paragraph(paragraph)}</li>')
    elif style == "Quote":
        close_list()
        body.append(f'<blockquote class="article-pullquote">{rich_paragraph(paragraph)}</blockquote>')
    else:
        close_list()
        body.append(f'<p>{rich_paragraph(paragraph)}</p>')
close_list()

if not lead:
    raise RuntimeError("Article lead not found")
if len(sources) < 18:
    raise RuntimeError(f"Too few documentary sources: {len(sources)}")

word_count = len(re.findall(r"\b[\wÀ-ÖØ-öø-ÿ’'-]+\b", " ".join(p.text for p in document.paragraphs)))
organization = {
    "@type": ["Organization", "NewsMediaOrganization"], "@id": SITE + "#organization",
    "name": "Uno Sguardo sull’Uomo", "alternateName": ["UNO SGUARDO SULL'UOMO", "unosguardosulluomo"],
    "url": SITE, "description": "Testata editoriale indipendente italiana di approfondimento: indagini, dati, documenti e ricostruzione dei fatti oltre la narrazione.",
    "email": "mailto:unosguardosulluomo@gmail.com", "logo": {"@type": "ImageObject", "url": SITE + "assets/testata.webp"},
    "sameAs": ["https://www.facebook.com/profile.php?id=61593043131759", "https://www.linkedin.com/company/unosguardosulluomo/"],
}
schema = {"@context": "https://schema.org", "@graph": [
    organization,
    {"@type": "WebSite", "@id": SITE + "#website", "url": SITE, "name": "Uno Sguardo sull’Uomo", "publisher": {"@id": SITE + "#organization"}, "inLanguage": "it-IT"},
    {"@type": "ImageObject", "@id": URL + "#primaryimage", "url": SITE + D["socialImage"], "contentUrl": SITE + D["socialImage"], "caption": D["title"], "width": D["socialImageWidth"], "height": D["socialImageHeight"], "inLanguage": "it-IT"},
    {"@type": "NewsArticle", "@id": URL + "#article", "headline": D["title"], "alternativeHeadline": D["deck"], "description": D["description"], "url": URL, "mainEntityOfPage": {"@id": URL + "#webpage"}, "datePublished": D["datePublished"], "dateModified": D["dateModified"], "image": [{"@id": URL + "#primaryimage"}], "thumbnailUrl": SITE + D["socialImage"], "inLanguage": "it-IT", "isAccessibleForFree": True, "articleSection": D["category"], "genre": "Dossier geopolitico", "keywords": D["tags"], "about": [{"@type": "Thing", "name": D["macroTopic"]}, {"@type": "Thing", "name": D["category"]}, *[{"@type": "Thing", "name": tag} for tag in D["tags"]]], "wordCount": word_count, "author": {"@type": "Organization", "name": D["schemaAuthor"], "url": SITE + "chi-siamo.html"}, "publisher": {"@id": SITE + "#organization"}, "isPartOf": {"@id": SITE + "#website"}},
    {"@type": "WebPage", "@id": URL + "#webpage", "url": URL, "name": D["title"], "description": D["description"], "inLanguage": "it-IT", "datePublished": D["datePublished"], "dateModified": D["dateModified"], "isPartOf": {"@id": SITE + "#website"}, "breadcrumb": {"@id": URL + "#breadcrumb"}, "primaryImageOfPage": {"@id": URL + "#primaryimage"}, "mainEntity": {"@id": URL + "#article"}},
    {"@type": "BreadcrumbList", "@id": URL + "#breadcrumb", "itemListElement": [{"@type": "ListItem", "position": 1, "name": "Prima pagina", "item": SITE}, {"@type": "ListItem", "position": 2, "name": D["category"], "item": SITE + D["categoryPath"]}, {"@type": "ListItem", "position": 3, "name": D["title"], "item": URL}]},
]}

page = f'''<!doctype html>
<html lang="it">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="robots" content="index, follow, max-image-preview:large">
  <meta name="description" content="{esc(D['description'])}">
  <meta property="og:site_name" content="UNO SGUARDO SULL'UOMO">
  <meta property="og:title" content="{esc(D['ogTitle'])}">
  <meta property="og:description" content="{esc(D['description'])}">
  <meta property="og:type" content="article"><meta property="og:url" content="{URL}">
  <meta property="og:image" content="{SITE+D['socialImage']}"><meta property="og:image:alt" content="{esc(D['imageAlt'])}">
  <meta property="og:image:width" content="{D['socialImageWidth']}"><meta property="og:image:height" content="{D['socialImageHeight']}">
  <meta name="twitter:card" content="summary_large_image"><meta name="twitter:title" content="{esc(D['ogTitle'])}">
  <meta name="twitter:description" content="{esc(D['description'])}"><meta name="twitter:image" content="{SITE+D['socialImage']}">
  <meta property="article:published_time" content="{D['datePublished']}"><meta property="article:modified_time" content="{D['dateModified']}">
  <meta property="article:section" content="{esc(D['category'])}"><meta name="author" content="{esc(D['schemaAuthor'])}">
  {''.join(f'<meta property="article:tag" content="{esc(tag)}">' for tag in D['tags'])}
  <link rel="canonical" href="{URL}"><link rel="alternate" type="application/rss+xml" title="Uno Sguardo sull’Uomo" href="{SITE}feed.xml">
  <title>{esc(D['seoTitle'])} — Uno Sguardo sull'Uomo</title>
  <link rel="stylesheet" href="styles.css?v=20261014-inline-image-1"><link rel="stylesheet" href="editorial-rules.css?v=20260812-1"><link rel="stylesheet" href="dossier-tables.css?v=20260925-1">
  <link rel="icon" type="image/svg+xml" href="assets/favicon.svg">
  <script type="application/ld+json" data-seo-schema>{json.dumps(schema, ensure_ascii=False, separators=(',', ':'))}</script>
</head>
<body><main class="newspaper">
  <div class="utility-bar"><span>Quotidiano indipendente di approfondimento</span><span>Politica, società e istituzioni</span><time class="current-date" datetime="{D['datePublished']}">{D['dateLabel']}</time></div>
  <header class="masthead"><a href="/" aria-label="Prima pagina"><img src="assets/testata.webp" width="1230" height="279" alt="Uno Sguardo sull’Uomo"></a></header>
  <nav class="nav" aria-label="Navigazione principale"><a href="/">Prima pagina</a><a href="indagini.html">Indagini</a><a href="ricerca-indagini.html">Ricerca indagini</a><a href="metodo.html">Metodo</a><a href="chi-siamo.html">Chi siamo</a><a href="contatti.html">Contatti</a></nav>
  <article class="article-page" data-macro-topic="{esc(D['macroTopic'])}" data-micro-topic="{esc(D['category'])}">
    <nav class="topic-breadcrumb" aria-label="Percorso tematico"><span class="topic-breadcrumb-macro">{esc(D['macroTopic'])}</span><span aria-hidden="true">›</span><a href="{D['categoryPath']}">{esc(D['category'])}</a><span aria-hidden="true">›</span><span aria-current="page">{esc(D['title'])}</span></nav>
    <header class="article-header"><a class="category-label" href="{D['categoryPath']}">{esc(D['category'])}</a><p class="eyebrow">DOSSIER</p><h1>{esc(D['title'])}</h1><p class="article-deck">{esc(D['deck'])}</p><div class="article-meta"><span>{esc(D['author'])}</span><time datetime="{D['datePublished']}">Pubblicato: {D['dateLabel']}</time><span>Tempo di lettura: {D['readingMinutes']} minuti</span></div><div class="topic-list" aria-label="Argomenti">{''.join('<span>'+esc(tag)+'</span>' for tag in D['tags'])}</div></header>
    <figure class="article-hero"><img src="{D['image']}" width="{D['imageWidth']}" height="{D['imageHeight']}" alt="{esc(D['imageAlt'])}" fetchpriority="high"><figcaption>{esc(D['imageCaption'])}</figcaption></figure>
    <div class="article-layout"><div class="article-body"><p class="article-lead">{lead}</p><details class="dossier-contents"><summary>In questo dossier</summary><ol>{''.join(contents)}<li><a href="#fonti">Fonti</a></li></ol></details>
{chr(10).join(body)}
<section class="sources" aria-labelledby="fonti"><h2 id="fonti">FONTI</h2><p>{source_intro}</p><ul class="sources-list">{''.join('<li>'+esc(source)+'</li>' for source in sources)}</ul><h3>Criterio editoriale</h3><p>{criterion}</p></section>
    </div></div>
  </article>
  <footer class="footer"><span>Uno Sguardo sull’Uomo — {esc(D['category'])}</span><span><a href="privacy-cookie.html">Privacy e Cookie Policy</a></span><span><a href="mailto:unosguardosulluomo@gmail.com">unosguardosulluomo@gmail.com</a></span></footer>
</main><script src="date.js?v=20260826-audit-1"></script></body></html>
'''
(ROOT / D["slug"]).write_text(page, encoding="utf-8", newline="\n")


def card(kind="archive-card", heading="h3", eager=False):
    priority = ' fetchpriority="high"' if eager else ' loading="lazy"'
    return (f'<article class="{kind}" data-dossier="{esc(D["slug"])}"><a class="story-image" href="{esc(D["slug"])}" aria-label="Leggi il dossier {esc(D["title"])}"><img src="{esc(D["image"])}" width="{D["imageWidth"]}" height="{D["imageHeight"]}" alt="{esc(D["imageAlt"])}"{priority}></a><p class="eyebrow">{esc(D["cardEyebrow"])}</p><p class="archive-meta"><time datetime="{D["datePublished"]}">{D["dateLabel"]}</time></p><{heading}><a class="headline-link" href="{esc(D["slug"])}">{esc(D["title"])}</a></{heading}><p>{esc(D["description"])}</p><a class="read-more" href="{esc(D["slug"])}">Leggi il dossier →</a></article>')


def dated_cards(markup):
    cards = re.findall(r'<article class="archive-card".*?</article>', markup, re.S)
    return sorted(cards, key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)


def change_card(markup, kind, heading):
    updated = re.sub(r'class="archive-card"', f'class="{kind}"', markup, count=1)
    updated = re.sub(r'<h3>', f'<{heading}>', updated, count=1)
    updated = re.sub(r'</h3>', f'</{heading}>', updated, count=1)
    return updated.replace(' fetchpriority="high"', ' loading="lazy"')


def update_current_date(markup):
    return re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', markup, count=1)


archive_path = ROOT / D["categoryPath"]
archive = archive_path.read_text(encoding="utf-8")
archive = re.sub(r'<article class="archive-card" data-dossier="' + re.escape(D["slug"]) + r'".*?</article>', "", archive, flags=re.S)
archive = re.sub(r'(<div class="archive-grid">)\s*', r'\1\n', archive, count=1)
archive = archive.replace('<div class="archive-grid">\n', '<div class="archive-grid">\n' + card() + '\n', 1)
archive = update_current_date(archive)
archive_path.write_text(archive, encoding="utf-8", newline="\n")

category_cards = dated_cards(archive)
indagini_path = ROOT / "indagini.html"
indagini = indagini_path.read_text(encoding="utf-8")
section = re.compile(r'(<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-politica-italiana.html">Politica italiana</a></h2></div><div class="archive-grid">).*?(</div><a class="category-archive-link" href="archivio-politica-italiana.html">)', re.S)
indagini, count = section.subn(lambda match: match.group(1) + "".join(category_cards[:3]) + match.group(2), indagini, count=1)
if count != 1:
    raise RuntimeError("Italian politics selection not found")
indagini = update_current_date(indagini)
indagini_path.write_text(indagini, encoding="utf-8", newline="\n")

all_archive_cards = []
for category_archive in ("archivio-politica-italiana.html", "archivio-politica-internazionale.html", "archivio-economia.html", "archivio-societa.html"):
    all_archive_cards.extend(dated_cards((ROOT / category_archive).read_text(encoding="utf-8")))
all_archive_cards.sort(key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)

home_path = ROOT / "index.html"
home = home_path.read_text(encoding="utf-8")
lead_card = change_card(all_archive_cards[0], "lead-story", "h1").replace(' loading="lazy"', ' fetchpriority="high"', 1)
side = "".join(change_card(item, "side-story", "h2") for item in all_archive_cards[1:3])
latest = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[3:6])
other = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[6:9])
home, count = re.subn(r'<article class="lead-story".*?</article>', lambda _: lead_card, home, count=1, flags=re.S)
if count != 1: raise RuntimeError("Home lead not found")
home, count = re.subn(r'(<aside>).*?(</aside>)', lambda match: match.group(1) + side + match.group(2), home, count=1, flags=re.S)
if count != 1: raise RuntimeError("Home sidebar not found")
home, count = re.subn(r'(<section class="story-grid" aria-label="Ultime indagini">).*?(</section>)', lambda match: match.group(1) + latest + match.group(2), home, count=1, flags=re.S)
if count != 1: raise RuntimeError("Home latest grid not found")
home, count = re.subn(r'(<section class="story-grid" aria-label="Altre indagini recenti">).*?(</section>)', lambda match: match.group(1) + other + match.group(2), home, count=1, flags=re.S)
if count != 1: raise RuntimeError("Home other grid not found")
home = update_current_date(home)
home_path.write_text(home, encoding="utf-8", newline="\n")

article_entry = f'  <url><loc>{URL}</loc><lastmod>{D["dateModified"]}</lastmod><image:image><image:loc>{SITE + D["socialImage"]}</image:loc></image:image></url>\n'
path = ROOT / "sitemap-articles.xml"
markup = re.sub(r'\s*<url><loc>' + re.escape(URL) + r'</loc>.*?</url>', "", path.read_text(encoding="utf-8"), flags=re.S)
markup = re.sub(r'(xmlns:image="http://www.google.com/schemas/sitemap-image/1.1">)\s*', r'\1\n' + article_entry, markup, count=1)
path.write_text(markup, encoding="utf-8", newline="\n")

news_entry = f'  <url><loc>{URL}</loc><news:news><news:publication><news:name>Uno Sguardo sull’Uomo</news:name><news:language>it</news:language></news:publication><news:publication_date>{D["datePublished"]}</news:publication_date><news:title>{esc(D["title"])}</news:title></news:news></url>\n'
path = ROOT / "sitemap-news.xml"
markup = re.sub(r'\s*<url><loc>' + re.escape(URL) + r'</loc>.*?</url>', "", path.read_text(encoding="utf-8"), flags=re.S)
markup = re.sub(r'(xmlns:news="http://www.google.com/schemas/sitemap-news/0.9">)\s*', r'\1\n' + news_entry, markup, count=1)
path.write_text(markup, encoding="utf-8", newline="\n")

pub_dt = datetime.fromisoformat(D["datePublished"]).replace(tzinfo=timezone.utc)
feed_item = f'''    <item>\n      <title>{esc(D['title'])}</title>\n      <link>{URL}</link>\n      <guid isPermaLink="true">{URL}</guid>\n      <pubDate>{email.utils.format_datetime(pub_dt)}</pubDate>\n      <category>{esc(D['category'])}</category>\n      <description>{esc(D['description'])}</description>\n    </item>\n'''
path = ROOT / "feed.xml"
markup = re.sub(r'\s*<item>\s*<title>' + re.escape(esc(D["title"])) + r'</title>.*?</item>', "", path.read_text(encoding="utf-8"), flags=re.S)
markup = markup.replace('    <atom:link href="https://unosguardosulluomo.github.io/feed.xml" rel="self" type="application/rss+xml"/>\n', '    <atom:link href="https://unosguardosulluomo.github.io/feed.xml" rel="self" type="application/rss+xml"/>\n' + feed_item, 1)
path.write_text(markup, encoding="utf-8", newline="\n")

for path_name in ("sitemap.xml", "sitemap-google.xml", "sitemap-index.xml"):
    path = ROOT / path_name
    path.write_text(re.sub(r'<lastmod>[^<]+</lastmod>', f'<lastmod>{D["dateModified"]}</lastmod>', path.read_text(encoding="utf-8")), encoding="utf-8", newline="\n")

print("Built Salvini-Lega dossier, editorial placements, feed and sitemaps")
