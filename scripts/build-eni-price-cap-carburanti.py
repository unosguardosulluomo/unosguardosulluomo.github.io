"""Build the Eni price-cap dossier and all editorial placements from the approved DOCX."""

import html
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/eni-price-cap-carburanti.json").read_text(encoding="utf-8"))
SITE = "https://unosguardosulluomo.github.io/"
URL = SITE + D["slug"]


def esc(value):
    return html.escape(str(value), quote=True)


def rich_paragraph(paragraph):
    fragments = []
    for run in paragraph.runs:
        value = esc(run.text).replace("\n", "<br>")
        if not value:
            continue
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments) or esc(paragraph.text)


def card(kind="archive-card", heading="h3", eager=False):
    priority = ' fetchpriority="high"' if eager else ' loading="lazy"'
    return (
        f'<article class="{kind}" data-dossier="{esc(D["slug"])}">'
        f'<a class="story-image" href="{esc(D["slug"])}" aria-label="Leggi il dossier {esc(D["title"])}">'
        f'<img src="{esc(D["image"])}" width="{D["imageWidth"]}" height="{D["imageHeight"]}" '
        f'alt="{esc(D["imageAlt"])}"{priority}></a>'
        f'<p class="eyebrow">{esc(D["cardEyebrow"])}</p>'
        f'<p class="archive-meta"><time datetime="{D["datePublished"]}">{D["dateLabel"]}</time></p>'
        f'<{heading}><a class="headline-link" href="{esc(D["slug"])}">{esc(D["title"])}</a></{heading}>'
        f'<p>{esc(D["description"])}</p>'
        f'<a class="read-more" href="{esc(D["slug"])}">Leggi il dossier →</a></article>'
    )


def inline_figure():
    return (
        f'<figure class="article-inline-figure"><img src="{esc(D["inlineImage"])}" '
        f'width="{D["inlineImageWidth"]}" height="{D["inlineImageHeight"]}" '
        f'alt="{esc(D["inlineImageAlt"])}" loading="lazy">'
        f'<figcaption>{esc(D["inlineImageCaption"])}</figcaption></figure>'
    )


def timeline(document):
    rows = document.tables[0].rows
    body = []
    for row in rows:
        body.append(
            "<tr>" + "".join(
                f'<th scope="row">{esc(cell.text)}</th>' if index == 0 else f"<td>{esc(cell.text)}</td>"
                for index, cell in enumerate(row.cells)
            ) + "</tr>"
        )
    return (
        '<p class="table-hint">Scorri orizzontalmente per leggere l’intera cronologia.</p>'
        '<div class="data-table-wrap" role="region" aria-label="Cronologia dell’indagine Eni" tabindex="0">'
        '<table class="data-table"><caption>Cronologia dell’indagine: 22–26 settembre 2026</caption>'
        f"<tbody>{''.join(body)}</tbody></table></div>"
    )


document = Document(ROOT / D["sourceDocx"])
standfirst = document.paragraphs[3].text.strip()
body = []
contents = []
sources = []
heading_number = 0
in_sources = False

for paragraph in document.paragraphs[5:]:
    text = paragraph.text.strip()
    if not text:
        continue
    style = paragraph.style.name
    if style == "Heading 1" and text == "FONTI":
        in_sources = True
        continue
    if style == "Heading 1" and text.startswith("SCHEDA TECNICA"):
        break
    if in_sources:
        sources.append(text)
        continue
    if style in ("Heading 1", "Heading 2"):
        if text == "IL NOME: CLAUDIO DESCALZI":
            body.append(inline_figure())
        heading_number += 1
        anchor = f"sezione-{heading_number}"
        contents.append(f'<li><a href="#{anchor}">{esc(text)}</a></li>')
        body.append(f'<h2 id="{anchor}">{esc(text)}</h2>')
    elif style == "Dateline":
        iso_date = {
            "22 settembre 2026": "2026-09-22",
            "25 settembre 2026": "2026-09-25",
            "26 settembre 2026": "2026-09-26",
        }.get(text, D["datePublished"])
        body.append(f'<p class="article-dateline"><time datetime="{iso_date}">{esc(text)}</time></p>')
    elif style == "Quote":
        body.append(f'<blockquote class="article-pullquote">{rich_paragraph(paragraph)}</blockquote>')
    else:
        body.append(f"<p>{rich_paragraph(paragraph)}</p>")

if not sources:
    raise RuntimeError("No document sources found")

organization = {
    "@type": ["Organization", "NewsMediaOrganization"],
    "@id": SITE + "#organization",
    "name": "Uno Sguardo sull’Uomo",
    "url": SITE,
    "logo": {"@type": "ImageObject", "url": SITE + "assets/testata.webp"},
}
schema = {
    "@context": "https://schema.org",
    "@graph": [
        organization,
        {
            "@type": "WebSite",
            "@id": SITE + "#website",
            "url": SITE,
            "name": "Uno Sguardo sull’Uomo",
            "publisher": {"@id": SITE + "#organization"},
            "inLanguage": "it-IT",
        },
        {
            "@type": "NewsArticle",
            "@id": URL + "#article",
            "headline": D["title"],
            "description": D["description"],
            "url": URL,
            "mainEntityOfPage": {"@type": "WebPage", "@id": URL},
            "datePublished": D["datePublished"],
            "dateModified": D["dateModified"],
            "image": [SITE + D["socialImage"]],
            "inLanguage": "it-IT",
            "articleSection": D["macroCategory"],
            "genre": "Dossier evolutivo",
            "about": [
                {"@type": "Thing", "name": D["macroCategory"]},
                {"@type": "Thing", "name": D["microCategory"]},
            ],
            "keywords": ", ".join(D["tags"]),
            "author": {"@type": "Organization", "name": "Redazione Uno Sguardo sull’Uomo", "url": SITE + "chi-siamo.html"},
            "publisher": {"@id": SITE + "#organization"},
            "isPartOf": {"@id": SITE + "#website"},
        },
        {
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"@type": "ListItem", "position": 1, "name": "Prima pagina", "item": SITE},
                {"@type": "ListItem", "position": 2, "name": D["macroCategory"], "item": SITE + D["categoryPath"]},
                {"@type": "ListItem", "position": 3, "name": D["microCategory"], "item": URL},
                {"@type": "ListItem", "position": 4, "name": D["title"], "item": URL},
            ],
        },
    ],
}

page = f'''<!doctype html>
<html lang="it">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="robots" content="index, follow, max-image-preview:large">
  <meta name="description" content="{esc(D['description'])}">
  <meta property="og:site_name" content="UNO SGUARDO SULL'UOMO">
  <meta property="og:title" content="{esc(D['title'])}">
  <meta property="og:description" content="{esc(D['description'])}">
  <meta property="og:type" content="article">
  <meta property="og:url" content="{URL}">
  <meta property="og:image" content="{SITE+D['socialImage']}">
  <meta property="og:image:alt" content="{esc(D['imageAlt'])}">
  <meta property="og:image:width" content="{D['socialImageWidth']}">
  <meta property="og:image:height" content="{D['socialImageHeight']}">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="{esc(D['title'])}">
  <meta name="twitter:description" content="{esc(D['description'])}">
  <meta name="twitter:image" content="{SITE+D['socialImage']}">
  <meta property="article:published_time" content="{D['datePublished']}">
  <meta property="article:modified_time" content="{D['dateModified']}">
  <meta property="article:section" content="{esc(D['macroCategory'])}">
  <meta property="article:tag" content="{esc(D['microCategory'])}">
  <link rel="canonical" href="{URL}">
  <title>Eni e price cap carburanti: controllo pubblico e Stato imprenditore — Uno Sguardo sull'Uomo</title>
  <link rel="stylesheet" href="styles.css?v=20260924-mobile-1">
  <link rel="stylesheet" href="editorial-rules.css?v=20260812-1">
  <link rel="stylesheet" href="dossier-tables.css?v=20260925-1">
  <link rel="icon" type="image/svg+xml" href="assets/favicon.svg">
  <script type="application/ld+json" data-seo-schema>{json.dumps(schema, ensure_ascii=False)}</script>
</head>
<body><main class="newspaper">
  <div class="utility-bar"><span>Quotidiano indipendente di approfondimento</span><span>Politica, società e istituzioni</span><time class="current-date" datetime="{D['datePublished']}">{D['dateLabel']}</time></div>
  <header class="masthead"><a href="/" aria-label="Prima pagina"><img src="assets/testata.webp" alt="Uno Sguardo sull’Uomo"></a></header>
  <nav class="nav" aria-label="Navigazione principale"><a href="/">Prima pagina</a><a href="indagini.html">Indagini</a><a href="metodo.html">Metodo</a><a href="chi-siamo.html">Chi siamo</a><a href="contatti.html">Contatti</a></nav>
  <article class="article-page" data-macro-category="{esc(D['macroCategory'])}" data-micro-category="{esc(D['microCategory'])}">
    <header class="article-header"><a class="category-label" href="{D['categoryPath']}">{esc(D['macroCategory'])}</a><p class="eyebrow">{esc(D['microCategory'])} · Dossier evolutivo</p><h1>{esc(D['title'])}</h1><p class="article-deck">{esc(D['deck'])}</p><div class="article-meta"><span>{esc(D['author'])}</span><time datetime="{D['dateOpened']}">Indagine aperta: {D['dateOpenedLabel']}</time><time datetime="{D['datePublished']}">Pubblicato: {D['dateLabel']}</time><time datetime="{D['dateModified']}">Aggiornato: {D['dateModifiedLabel']}</time><span>Tempo di lettura: {D['readingMinutes']} minuti</span></div><div class="topic-list" aria-label="Argomenti">{''.join('<span>'+esc(tag)+'</span>' for tag in D['tags'])}</div></header>
    <figure class="article-hero"><img src="{D['image']}" width="{D['imageWidth']}" height="{D['imageHeight']}" alt="{esc(D['imageAlt'])}" fetchpriority="high"><figcaption>{esc(D['imageCaption'])}</figcaption></figure>
    <div class="article-layout"><div class="article-body"><p class="standfirst">{esc(standfirst)}</p>{timeline(document)}<details class="dossier-contents"><summary>In questo dossier</summary><ol>{''.join(contents)}<li><a href="#fonti">Fonti documentali</a></li></ol></details>
{chr(10).join(body)}
<aside class="related-dossiers" aria-labelledby="approfondimenti"><h3 id="approfondimenti">Approfondimenti collegati</h3><ul><li><a href="article-eni-baleine.html">Eni, Baleine e la politica energetica italiana</a></li><li><a href="article-cdp.html">Cassa Depositi e Prestiti: lo Stato azionista</a></li></ul></aside>
<section class="sources" aria-labelledby="fonti"><h2 id="fonti">FONTI DOCUMENTALI</h2><ul class="sources-list">{''.join('<li>'+esc(source)+'</li>' for source in sources)}</ul></section>
    </div></div>
  </article>
  <footer class="footer"><span>Uno Sguardo sull’Uomo — {esc(D['macroCategory'])}</span><span><a href="privacy-cookie.html">Privacy e Cookie Policy</a></span><span><a href="mailto:unosguardosulluomo@gmail.com">unosguardosulluomo@gmail.com</a></span></footer>
</main><script src="seo.js?v=20260903-google-1"></script><script src="date.js?v=20260826-audit-1"></script></body></html>
'''
(ROOT / D["slug"]).write_text(page, encoding="utf-8")


def dated_cards(markup):
    cards = re.findall(r'<article class="archive-card".*?</article>', markup, re.S)
    return sorted(cards, key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)


def change_card(card_markup, kind, heading):
    updated = re.sub(r'class="archive-card"', f'class="{kind}"', card_markup, count=1)
    updated = re.sub(r'<h3>', f'<{heading}>', updated, count=1)
    updated = re.sub(r'</h3>', f'</{heading}>', updated, count=1)
    updated = updated.replace(' fetchpriority="high"', ' loading="lazy"')
    if ' loading="lazy"' not in updated:
        updated = updated.replace('></a><p class="eyebrow">', ' loading="lazy"></a><p class="eyebrow">', 1)
    return updated


archive_path = ROOT / D["categoryPath"]
archive = archive_path.read_text(encoding="utf-8")
archive = re.sub(r'<article class="archive-card" data-dossier="' + re.escape(D["slug"]) + r'".*?</article>', "", archive, flags=re.S)
archive = re.sub(r'(<div class="archive-grid">)\s*', r'\1\n', archive, count=1)
archive = archive.replace('<div class="archive-grid">\n', '<div class="archive-grid">\n' + card() + '\n', 1)
archive = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', archive, count=1)
archive_path.write_text(archive, encoding="utf-8")

category_cards = dated_cards(archive)
indagini_path = ROOT / "indagini.html"
indagini = indagini_path.read_text(encoding="utf-8")
section = re.compile(r'(<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-economia.html">Economia</a></h2></div><div class="archive-grid">).*?(</div><a class="category-archive-link" href="archivio-economia.html">)', re.S)
indagini, count = section.subn(lambda match: match.group(1) + "".join(category_cards[:3]) + match.group(2), indagini, count=1)
assert count == 1, "Economia selection not found"
indagini = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', indagini, count=1)
indagini_path.write_text(indagini, encoding="utf-8")

all_archive_cards = []
for category_archive in ("archivio-politica-italiana.html", "archivio-politica-internazionale.html", "archivio-economia.html", "archivio-societa.html"):
    all_archive_cards.extend(dated_cards((ROOT / category_archive).read_text(encoding="utf-8")))
all_archive_cards = sorted(all_archive_cards, key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)

home_path = ROOT / "index.html"
home = home_path.read_text(encoding="utf-8")
lead = change_card(all_archive_cards[0], "lead-story", "h1").replace(' loading="lazy"', ' fetchpriority="high"', 1)
side_cards = "".join(change_card(item, "side-story", "h2") for item in all_archive_cards[1:3])
latest_cards = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[3:6])
other_cards = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[6:9])
home, count = re.subn(r'<article class="lead-story".*?</article>', lambda _: lead, home, count=1, flags=re.S)
assert count == 1, "Home lead not found"
home, count = re.subn(r'(<aside>).*?(</aside>)', lambda match: match.group(1) + side_cards + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home sidebar not found"
home, count = re.subn(r'(<section class="story-grid" aria-label="Ultime indagini">).*?(</section>)', lambda match: match.group(1) + latest_cards + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home latest grid not found"
home, count = re.subn(r'(<section class="story-grid" aria-label="Altre indagini recenti">).*?(</section>)', lambda match: match.group(1) + other_cards + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home other grid not found"
home = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', home, count=1)
home_path.write_text(home, encoding="utf-8")

for filename in ("sitemap.xml", "sitemap-google.xml"):
    path = ROOT / filename
    xml = path.read_text(encoding="utf-8")
    for suffix in ("", "indagini.html", D["categoryPath"], D["slug"]):
        location = SITE + suffix
        entry = f'<url><loc>{location}</loc><lastmod>{D["datePublished"]}</lastmod></url>'
        pattern = re.compile(r'<url>\s*<loc>' + re.escape(location) + r'</loc>.*?</url>', re.S)
        xml = pattern.sub(entry, xml, count=1) if pattern.search(xml) else xml.replace('</urlset>', '  ' + entry + '\n</urlset>')
    path.write_text(xml, encoding="utf-8")

print("Built Eni price-cap dossier, placements and sitemaps")
