"""Validate Eni dossier preservation, SEO metadata, assets, links and placements."""

import html
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/eni-price-cap-carburanti.json").read_text(encoding="utf-8"))
page = (ROOT / D["slug"]).read_text(encoding="utf-8")


def normalize(value):
    return re.sub(r"\s+", " ", value).strip()


assert D["slug"] == "article-economia-energia-eni-price-cap-carburanti.html"
assert "economia" in D["slug"] and "energia" in D["slug"]
assert f'<link rel="canonical" href="https://unosguardosulluomo.github.io/{D["slug"]}">' in page
assert f'<meta property="og:url" content="https://unosguardosulluomo.github.io/{D["slug"]}">' in page
assert f'<meta property="article:published_time" content="{D["datePublished"]}">' in page
assert f'<meta property="article:modified_time" content="{D["dateModified"]}">' in page
assert f'data-macro-category="{D["macroCategory"]}"' in page
assert f'data-micro-category="{D["microCategory"]}"' in page
assert page.count('<figure class="') == 2
assert page.count("intelligenza artificiale") == 2
assert page.count('<table class="data-table">') == 1
assert page.count('<blockquote class="article-pullquote">') == 1
assert "SCHEDA TECNICA PER CODEX" not in page and "ISTRUZIONE A CODEX" not in page

sources_html = re.search(r'<ul class="sources-list">(.*?)</ul>', page, re.S).group(1)
assert sources_html.count("<li>") == 12
assert "<a " not in sources_html
assert "http://" not in sources_html and "https://" not in sources_html

plain = normalize(html.unescape(re.sub(r"<[^>]+>", " ", page)))
document = Document(ROOT / D["sourceDocx"])
missing = []
for paragraph in [document.paragraphs[3], *document.paragraphs[5:73]]:
    value = normalize(paragraph.text)
    if value and value != "FONTI" and value not in plain:
        missing.append(value)
assert not missing, "Missing DOCX paragraphs: " + " | ".join(missing[:5])
for row in document.tables[0].rows:
    for cell in row.cells:
        assert normalize(cell.text) in plain, f"Missing timeline text: {cell.text}"

for expanded in (
    "Ministero dell’Economia e delle Finanze (MEF)",
    "Cassa Depositi e Prestiti (CDP)",
):
    assert expanded in plain, expanded

for filename in ("index.html", "indagini.html", "archivio-economia.html", "sitemap.xml", "sitemap-google.xml"):
    assert D["slug"] in (ROOT / filename).read_text(encoding="utf-8"), f"Missing placement in {filename}"

home = (ROOT / "index.html").read_text(encoding="utf-8")
assert D["slug"] in re.search(r'<article class="lead-story".*?</article>', home, re.S).group(0)

indagini = (ROOT / "indagini.html").read_text(encoding="utf-8")
economy = re.search(r'<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-economia.html">Economia</a></h2></div><div class="archive-grid">(.*?)</div><a class="category-archive-link" href="archivio-economia.html">', indagini, re.S).group(1)
assert economy.count('<article class="archive-card"') == 3
assert economy.index(D["slug"]) < economy.index("article-giorgetti-pallottoliere.html")

archive = (ROOT / "archivio-economia.html").read_text(encoding="utf-8")
archive_grid = re.search(r'<div class="archive-grid">(.*?)</div><footer', archive, re.S).group(1)
assert archive_grid.index(D["slug"]) < archive_grid.index("article-giorgetti-pallottoliere.html")

for value in re.findall(r'(?:href|src)="([^"]+)"', page):
    if not value or value.startswith(("http://", "https://", "mailto:", "#", "/")):
        continue
    assert (ROOT / value.split("?", 1)[0]).exists(), f"Broken local reference: {value}"

schema = json.loads(re.search(r'<script type="application/ld\+json" data-seo-schema>(.*?)</script>', page, re.S).group(1))
article = next(item for item in schema["@graph"] if item.get("@type") == "NewsArticle")
assert article["datePublished"] == D["datePublished"]
assert article["dateModified"] == D["dateModified"]
assert article["articleSection"] == D["macroCategory"]
assert [item["name"] for item in article["about"]] == [D["macroCategory"], D["microCategory"]]
assert article["image"] == ["https://unosguardosulluomo.github.io/" + D["socialImage"]]

for sitemap in ("sitemap.xml", "sitemap-google.xml"):
    ET.parse(ROOT / sitemap)

print("Eni price-cap dossier checks passed")
