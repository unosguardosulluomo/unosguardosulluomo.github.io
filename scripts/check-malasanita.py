"""Validate the Malasanità publication, preservation, metadata and placements."""

import html
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/malasanita.json").read_text(encoding="utf-8"))
page = (ROOT / D["slug"]).read_text(encoding="utf-8")


def normalize(value):
    return re.sub(r"\s+", " ", value).strip()


assert f'<link rel="canonical" href="https://unosguardosulluomo.github.io/{D["slug"]}">' in page
assert f'<meta property="og:url" content="https://unosguardosulluomo.github.io/{D["slug"]}">' in page
assert f'<meta property="article:published_time" content="{D["datePublished"]}">' in page
assert f'<meta property="article:section" content="{D["category"]}">' in page
assert page.count('<figure class="') == 2
assert page.count("intelligenza artificiale") == 2
assert page.count('<table class="data-table">') == 1
assert page.count('<blockquote class="article-testimony">') == 1
assert page.count('<blockquote class="article-pullquote">') >= 7

sources_html = re.search(r'<ul class="sources-list">(.*?)</ul>', page, re.S).group(1)
assert sources_html.count("<li>") == 22
assert "<a " not in sources_html
assert "http://" not in sources_html and "https://" not in sources_html

plain = normalize(html.unescape(re.sub(r"<[^>]+>", " ", page)))
document = Document(ROOT / D["sourceDocx"])
missing = []
for paragraph in [document.paragraphs[3], *document.paragraphs[5:]]:
    text = normalize(paragraph.text)
    if not text:
        continue
    if text.startswith(("•", "", "�")):
        text = normalize(text[1:])
    if text not in plain:
        missing.append(text)
assert not missing, "Missing DOCX paragraphs: " + " | ".join(missing[:5])

for table in document.tables[1:]:
    for row in table.rows:
        for cell in row.cells:
            for line in cell.text.splitlines():
                value = normalize(line)
                assert not value or value in plain, f"Missing table text: {value}"

for expanded in (
    "Servizio sanitario nazionale (SSN)",
    "Agenzia di tutela della salute (ATS)",
    "Servizio sanitario regionale (SSR)",
    "Responsabile unico aziendale per i tempi d’attesa (RUA)",
    "Ufficio relazioni con il pubblico (URP)",
    "attività libero-professionale intramuraria (ALPI)",
    "Organismo regionale per le attività di controllo (ORAC)",
    "Piano nazionale di ripresa e resilienza (PNRR)",
):
    assert expanded in plain, expanded

for filename in ("index.html", "indagini.html", "archivio-societa.html", "sitemap.xml", "sitemap-google.xml"):
    value = (ROOT / filename).read_text(encoding="utf-8")
    assert D["slug"] in value, f"Missing placement in {filename}"

home = (ROOT / "index.html").read_text(encoding="utf-8")
lead = re.search(r'<article class="lead-story".*?</article>', home, re.S).group(0)
assert D["slug"] in lead

indagini = (ROOT / "indagini.html").read_text(encoding="utf-8")
society_section = re.search(
    r'<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-societa.html">Società</a></h2></div><div class="archive-grid">(.*?)</div><a class="category-archive-link" href="archivio-societa.html">',
    indagini,
    re.S,
).group(1)
assert society_section.count('<article class="archive-card"') == 3
assert society_section.index(D["slug"]) < society_section.index("article-i-famosi-bulli.html")

archive = (ROOT / "archivio-societa.html").read_text(encoding="utf-8")
archive_grid = re.search(r'<div class="archive-grid">(.*?)</div><footer', archive, re.S).group(1)
assert archive_grid.index(D["slug"]) < archive_grid.index("article-i-famosi-bulli.html")

for value in re.findall(r'(?:href|src)="([^"]+)"', page):
    if not value or value.startswith(("http://", "https://", "mailto:", "#", "/")):
        continue
    target = ROOT / value.split("?", 1)[0]
    assert target.exists(), f"Broken local reference: {value}"

schema_text = re.search(r'<script type="application/ld\+json" data-seo-schema>(.*?)</script>', page, re.S).group(1)
schema = json.loads(schema_text)
article = next(item for item in schema["@graph"] if item.get("@type") == "NewsArticle")
assert article["datePublished"] == D["datePublished"]
assert article["articleSection"] == D["category"]
assert article["image"] == ["https://unosguardosulluomo.github.io/" + D["socialImage"]]

print("Malasanità dossier checks passed")
