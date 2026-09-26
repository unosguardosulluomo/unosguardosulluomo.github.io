from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "dossier-eni-price-cap-carburanti.docx"


def replace_paragraph(document, starts_with, replacement):
    matches = [p for p in document.paragraphs if p.text.startswith(starts_with)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {starts_with!r}, found {len(matches)}")
    paragraph = matches[0]
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(replacement)


document = Document(PATH)

replace_paragraph(
    document,
    "Il Ministero dell’Economia e delle Finanze possiede",
    "Il Ministero dell’Economia e delle Finanze (MEF) possiede direttamente il 2,17% del capitale; "
    "Cassa Depositi e Prestiti (CDP) il 30,92%. La partecipazione complessiva riconducibile al settore "
    "pubblico arriva così al 33,09%. Sul proprio sito istituzionale Eni indica che il MEF esercita il "
    "“controllo di fatto” sulla società attraverso la quota diretta e quella detenuta tramite CDP.",
)

replace_paragraph(
    document,
    "Questo non significa che Palazzo Chigi possa fissare",
    "Eni non è più un ente pubblico: è una società per azioni quotata. Ma il MEF ne esercita il controllo "
    "di fatto attraverso la partecipazione diretta e quella detenuta tramite CDP. In questo senso lo Stato "
    "italiano continua a fare impresa nel settore energetico come azionista di controllo, senza per questo "
    "gestire direttamente le singole decisioni commerciali della società.",
)

replace_paragraph(
    document,
    "Quello che il 22 settembre era il centro della nostra domanda",
    "La decisione di Eni è coerente con l’ipotesi formulata il 22 settembre: usare il prezzo come leva "
    "competitiva. Non dimostra da sola che il Governo l’abbia richiesta o determinata, ma mostra che un "
    "grande operatore controllato di fatto dallo Stato può intervenire sul proprio prezzo per esercitare "
    "pressione concorrenziale.",
)

replace_paragraph(
    document,
    "Il 25 settembre Eni ha annunciato esattamente",
    "Il 25 settembre Eni ha annunciato l’ingresso su quel terreno: non un prezzo imposto a tutti, ma un "
    "tetto massimo sulla propria rete per trenta giorni. La decisione appartiene alla società; non prova, "
    "in assenza di documenti, che sia stata ordinata dal Governo.",
)

document.save(PATH)
print(f"Revised {PATH.name}")
