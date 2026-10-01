"""Build the Ravenna CPR dossier and all editorial placements from the approved DOCX."""

import html
import json
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/medici-sotto-torchio.json").read_text(encoding="utf-8"))
SITE = "https://unosguardosulluomo.github.io/"
URL = SITE + D["slug"]


def esc(value):
    return html.escape(str(value), quote=True)


def rich_paragraph(paragraph):
    fragments = []
    for run in paragraph.runs:
        value = esc(run.text).replace("\n", "<br>")
        if not value:
            continue
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments) or esc(paragraph.text)


def blocks(parent):
    if not isinstance(parent, DocumentObject):
        raise TypeError("The dossier source must be a DOCX document")
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def card(kind="archive-card", heading="h3", eager=False):
    priority = ' fetchpriority="high"' if eager else ' loading="lazy"'
    return (
        f'<article class="{kind}" data-dossier="{esc(D["slug"])}">'
        f'<a class="story-image" href="{esc(D["slug"])}" aria-label="Leggi il dossier {esc(D["title"])}">'
        f'<img src="{esc(D["image"])}" width="{D["imageWidth"]}" height="{D["imageHeight"]}" '
        f'alt="{esc(D["imageAlt"])}"{priority}></a>'
        f'<p class="eyebrow">{esc(D["cardEyebrow"])}</p>'
        f'<p class="archive-meta"><time datetime="{D["datePublished"]}">{D["dateLabel"]}</time></p>'
        f'<{heading}><a class="headline-link" href="{esc(D["slug"])}">{esc(D["title"])}</a></{heading}>'
        f'<p>{esc(D["description"])}</p>'
        f'<a class="read-more" href="{esc(D["slug"])}">Leggi il dossier →</a></article>'
    )


def data_table(table):
    rows = table.rows
    head = "".join(f'<th scope="col">{esc(cell.text)}</th>' for cell in rows[0].cells)
    body = []
    for row in rows[1:]:
        cells = []
        for index, cell in enumerate(row.cells):
            tag = "th" if index == 0 else "td"
            scope = ' scope="row"' if index == 0 else ""
            cells.append(f"<{tag}{scope}>{esc(cell.text)}</{tag}>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return (
        '<p class="table-hint">Scorri orizzontalmente per vedere tutte le colonne.</p>'
        '<div class="data-table-wrap" role="region" aria-label="Esiti delle visite preventive a Ravenna" tabindex="0">'
        '<table class="data-table"><caption>Esiti delle 64 visite preventive esaminate nell’inchiesta di Ravenna</caption>'
        f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></div>"
    )


document = Document(ROOT / D["sourceDocx"])
standfirst = document.paragraphs[13].text.strip()
opening_quote = document.paragraphs[12].text.strip()
body = []
contents = []
sources = []
in_article = False
in_sources = False
heading_number = 0

for item in blocks(document):
    if isinstance(item, Table):
        if in_article and not in_sources:
            body.append(data_table(item))
        continue

    text = item.text.strip()
    if not text:
        continue
    style = item.style.name
    if not in_article:
        if style == "Heading 1" and text == "Il punto di partenza":
            in_article = True
        else:
            continue
    if style == "Heading 1" and text == "Fonti documentali":
        in_sources = True
        continue
    if in_sources:
        sources.append(text)
        continue
    if style == "Heading 1":
        heading_number += 1
        anchor = f"sezione-{heading_number}"
        contents.append(f'<li><a href="#{anchor}">{esc(text)}</a></li>')
        body.append(f'<h2 id="{anchor}">{esc(text)}</h2>')
    elif style == "Heading 2":
        body.append(f"<h3>{esc(text)}</h3>")
    elif style == "Richiamo":
        body.append(f'<blockquote class="article-pullquote">{rich_paragraph(item)}</blockquote>')
    else:
        body.append(f"<p>{rich_paragraph(item)}</p>")

if not sources:
    raise RuntimeError("No documentary sources found in the DOCX")

organization = {
    "@type": ["Organization", "NewsMediaOrganization"],
    "@id": SITE + "#organization",
    "name": "Uno Sguardo sull’Uomo",
    "url": SITE,
    "logo": {"@type": "ImageObject", "url": SITE + "assets/testata.webp"},
}
schema = {
    "@context": "https://schema.org",
    "@graph": [
        organization,
        {"@type": "WebSite", "@id": SITE + "#website", "url": SITE, "name": "Uno Sguardo sull’Uomo", "publisher": {"@id": SITE + "#organization"}, "inLanguage": "it-IT"},
        {
            "@type": "NewsArticle",
            "@id": URL + "#article",
            "headline": D["title"],
            "description": D["description"],
            "url": URL,
            "mainEntityOfPage": {"@type": "WebPage", "@id": URL},
            "datePublished": D["datePublished"],
            "image": [SITE + D["socialImage"]],
            "inLanguage": "it-IT",
            "articleSection": D["category"],
            "keywords": ", ".join(D["tags"]),
            "author": {"@type": "Organization", "name": "Redazione Uno Sguardo sull’Uomo", "url": SITE + "chi-siamo.html"},
            "publisher": {"@id": SITE + "#organization"},
            "isPartOf": {"@id": SITE + "#website"},
        },
        {
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"@type": "ListItem", "position": 1, "name": "Prima pagina", "item": SITE},
                {"@type": "ListItem", "position": 2, "name": D["category"], "item": SITE + D["categoryPath"]},
                {"@type": "ListItem", "position": 3, "name": D["title"], "item": URL},
            ],
        },
    ],
}

page = f'''<!doctype html>
<html lang="it">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="robots" content="index, follow, max-image-preview:large">
  <meta name="description" content="{esc(D['description'])}">
  <meta property="og:site_name" content="UNO SGUARDO SULL'UOMO">
  <meta property="og:title" content="{esc(D['title'])}">
  <meta property="og:description" content="{esc(D['description'])}">
  <meta property="og:type" content="article">
  <meta property="og:url" content="{URL}">
  <meta property="og:image" content="{SITE+D['socialImage']}">
  <meta property="og:image:alt" content="{esc(D['imageAlt'])}">
  <meta property="og:image:width" content="{D['socialImageWidth']}">
  <meta property="og:image:height" content="{D['socialImageHeight']}">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="{esc(D['title'])}">
  <meta name="twitter:description" content="{esc(D['description'])}">
  <meta name="twitter:image" content="{SITE+D['socialImage']}">
  <meta property="article:published_time" content="{D['datePublished']}">
  <meta property="article:section" content="{D['category']}">
  <link rel="canonical" href="{URL}">
  <title>{esc(D['title'])} — Uno Sguardo sull'Uomo</title>
  <link rel="stylesheet" href="styles.css?v=20260924-mobile-1">
  <link rel="stylesheet" href="editorial-rules.css?v=20260812-1">
  <link rel="stylesheet" href="dossier-tables.css?v=20260925-1">
  <link rel="icon" type="image/svg+xml" href="assets/favicon.svg">
  <script type="application/ld+json" data-seo-schema>{json.dumps(schema, ensure_ascii=False)}</script>
</head>
<body><main class="newspaper">
  <div class="utility-bar"><span>Quotidiano indipendente di approfondimento</span><span>Politica, società e istituzioni</span><time class="current-date" datetime="{D['datePublished']}">{D['dateLabel']}</time></div>
  <header class="masthead"><a href="/" aria-label="Prima pagina"><img src="assets/testata.webp" alt="Uno Sguardo sull’Uomo"></a></header>
  <nav class="nav" aria-label="Navigazione principale"><a href="/">Prima pagina</a><a href="indagini.html">Indagini</a><a href="metodo.html">Metodo</a><a href="chi-siamo.html">Chi siamo</a><a href="contatti.html">Contatti</a></nav>
  <article class="article-page">
    <header class="article-header"><a class="category-label" href="{D['categoryPath']}">{D['category']}</a><p class="eyebrow">DOSSIER</p><h1>{esc(D['title'])}</h1><p class="article-deck">{esc(D['deck'])}</p><div class="article-meta"><span>{D['author']}</span><time datetime="{D['datePublished']}">Pubblicato: {D['dateLabel']}</time><span>Tempo di lettura: {D['readingMinutes']} minuti</span></div><div class="topic-list" aria-label="Argomenti">{''.join('<span>'+esc(tag)+'</span>' for tag in D['tags'])}</div></header>
    <figure class="article-hero"><img src="{D['image']}" width="{D['imageWidth']}" height="{D['imageHeight']}" alt="{esc(D['imageAlt'])}" fetchpriority="high"><figcaption>{esc(D['imageCaption'])}</figcaption></figure>
    <div class="article-layout"><div class="article-body"><blockquote class="article-pullquote">{esc(opening_quote)}</blockquote><p class="standfirst">{esc(standfirst)}</p><details class="dossier-contents"><summary>In questo dossier</summary><ol>{''.join(contents)}<li><a href="#fonti">Fonti documentali</a></li></ol></details>
{chr(10).join(body)}
<section class="sources" aria-labelledby="fonti"><h2 id="fonti">FONTI DOCUMENTALI</h2><ul class="sources-list">{''.join('<li>'+esc(source)+'</li>' for source in sources)}</ul></section>
    </div></div>
  </article>
  <footer class="footer"><span>Uno Sguardo sull’Uomo — {D['category']}</span><span><a href="privacy-cookie.html">Privacy e Cookie Policy</a></span><span><a href="mailto:unosguardosulluomo@gmail.com">unosguardosulluomo@gmail.com</a></span></footer>
</main><script src="date.js?v=20260826-audit-1"></script></body></html>
'''
(ROOT / D["slug"]).write_text(page, encoding="utf-8")


def dated_cards(markup):
    cards = re.findall(r'<article class="archive-card".*?</article>', markup, re.S)
    return sorted(cards, key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)


def change_card(markup, kind, heading):
    updated = re.sub(r'class="archive-card"', f'class="{kind}"', markup, count=1)
    updated = re.sub(r'<h3>', f'<{heading}>', updated, count=1)
    updated = re.sub(r'</h3>', f'</{heading}>', updated, count=1)
    updated = updated.replace(' fetchpriority="high"', ' loading="lazy"')
    return updated


archive_path = ROOT / D["categoryPath"]
archive = archive_path.read_text(encoding="utf-8")
for managed_slug in [D["slug"], *D.get("legacySlugs", [])]:
    archive = re.sub(r'<article class="archive-card" data-dossier="' + re.escape(managed_slug) + r'".*?</article>', "", archive, flags=re.S)
archive = re.sub(r'(<div class="archive-grid">)\s*', r'\1\n', archive, count=1)
archive = archive.replace('<div class="archive-grid">\n', '<div class="archive-grid">\n' + card() + '\n', 1)
archive = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', archive, count=1)
archive_path.write_text(archive, encoding="utf-8")

category_cards = dated_cards(archive)
indagini_path = ROOT / "indagini.html"
indagini = indagini_path.read_text(encoding="utf-8")
section = re.compile(r'(<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-politica-italiana.html">Politica italiana</a></h2></div><div class="archive-grid">).*?(</div><a class="category-archive-link" href="archivio-politica-italiana.html">)', re.S)
indagini, count = section.subn(lambda match: match.group(1) + "".join(category_cards[:3]) + match.group(2), indagini, count=1)
assert count == 1, "Politica italiana selection not found"
indagini = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', indagini, count=1)
indagini_path.write_text(indagini, encoding="utf-8")

all_archive_cards = []
for category_archive in ("archivio-politica-italiana.html", "archivio-politica-internazionale.html", "archivio-economia.html", "archivio-societa.html"):
    all_archive_cards.extend(dated_cards((ROOT / category_archive).read_text(encoding="utf-8")))
all_archive_cards.sort(key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)

home_path = ROOT / "index.html"
home = home_path.read_text(encoding="utf-8")
lead = change_card(all_archive_cards[0], "lead-story", "h1").replace(' loading="lazy"', ' fetchpriority="high"', 1)
side = "".join(change_card(item, "side-story", "h2") for item in all_archive_cards[1:3])
latest = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[3:6])
other = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[6:9])
home, count = re.subn(r'<article class="lead-story".*?</article>', lambda _: lead, home, count=1, flags=re.S)
assert count == 1, "Home lead not found"
home, count = re.subn(r'(<aside>).*?(</aside>)', lambda match: match.group(1) + side + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home sidebar not found"
home, count = re.subn(r'(<section class="story-grid" aria-label="Ultime indagini">).*?(</section>)', lambda match: match.group(1) + latest + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home latest grid not found"
home, count = re.subn(r'(<section class="story-grid" aria-label="Altre indagini recenti">).*?(</section>)', lambda match: match.group(1) + other + match.group(2), home, count=1, flags=re.S)
assert count == 1, "Home other grid not found"
home = re.sub(r'<time class="current-date" datetime="[^"]+">.*?</time>', f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>', home, count=1)
home_path.write_text(home, encoding="utf-8")

print("Built Medici sotto torchio dossier and editorial placements")
