from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "dossier-trump-ricucire-giorgia.docx"


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def insert_before(document, reference, text, style):
    paragraph = document.add_paragraph(text, style=style)
    reference._p.addprevious(paragraph._p)
    return paragraph


document = Document(DOCX)

replacements = {
    "Il 19 giugno, dopo il G7,": "Il 19 giugno, dopo il Gruppo dei Sette (G7),",
    "conferenza economica Italia-USA": "conferenza economica Italia-Stati Uniti d’America (USA)",
    "OCSIT detiene": "L’Organismo centrale di stoccaggio italiano (OCSIT) detiene",
    "nel GNL e nell’approvvigionamento": "nel gas naturale liquefatto (GNL) e nell’approvvigionamento",
    "Il Ministero dell’Economia, direttamente e attraverso Cassa Depositi e Prestiti,": (
        "Il Ministero dell’economia e delle finanze (MEF), direttamente e attraverso "
        "Cassa depositi e prestiti (CDP),"
    ),
    "CDP, SACE, Ferrovie": "CDP, Servizi assicurativi del commercio estero (SACE), Ferrovie",
}

in_sources = False
for paragraph in document.paragraphs:
    if paragraph.text.strip() == "Fonti":
        in_sources = True
    if paragraph.style.name == "Title":
        set_paragraph_text(paragraph, "TRUMP SI INGINOCCHIA E CERCA DI RICUCIRE CON GIORGIA")
        continue
    if paragraph.style.name == "Deck":
        set_paragraph_text(
            paragraph,
            "Diesel, Eni e il potere italiano che Washington aveva sottovalutato",
        )
        continue
    if in_sources:
        continue
    text = paragraph.text.replace("ENI", "Eni")
    for old, new in replacements.items():
        text = text.replace(old, new)
    if text != paragraph.text:
        set_paragraph_text(paragraph, text)

target_heading = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.startswith("Poi il tavolo si gira")
)

new_york = [
    (
        "Il segnale di New York: Meloni non offre una ricucitura gratuita",
        "Heading 1",
    ),
    (
        "Alla fine di settembre Meloni manda un altro segnale. Non partecipa all’Assemblea "
        "generale dell’Organizzazione delle Nazioni Unite (ONU) a New York, dove Trump è "
        "presente e incontra diversi leader, e delega la rappresentanza italiana al ministro "
        "degli Esteri Antonio Tajani. Negli stessi giorni partecipa alla cena inaugurale della "
        "Milano Fashion Week.",
        "Normal",
    ),
    (
        "La spiegazione ufficiale è che l’Italia è pienamente rappresentata da Tajani e che la "
        "presidente del Consiglio deve seguire anche i dossier interni. Meloni rivendica inoltre "
        "la scelta di essere a Milano come segnale di vicinanza a una filiera che vale oltre 36 "
        "miliardi di euro di esportazioni. Tajani esclude pubblicamente che l’assenza dipenda "
        "dalle frizioni con Trump.",
        "Normal",
    ),
    (
        "Non è la prova di un ultimatum personale e non dimostra che Meloni abbia preteso delle "
        "scuse. Ma, dopo la derisione pubblica e lo scontro sulle basi, la scelta assume anche un "
        "significato politico: Meloni evita di offrire a Trump una ricucitura gratuita e "
        "pubblicamente visibile. Pochi giorni dopo, quando Washington ha bisogno di una risposta "
        "europea sul diesel, l’Italia torna al tavolo da una posizione diversa.",
        "Normal",
    ),
]
for text, style in new_york:
    insert_before(document, target_heading, text, style)

price_cap = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.startswith("Il 25 settembre Eni annuncia un tetto")
)
price_cap_text = price_cap.text.replace(
    "Nella stessa comunicazione Eni collega esplicitamente",
    "Eni precisa inoltre che il price cap è collegato all’agevolazione fiscale sulle accise "
    "in vigore. Nella stessa comunicazione Eni collega esplicitamente",
)
set_paragraph_text(price_cap, price_cap_text)

sources = [
    "Adnkronos, 4 settembre 2026 — «Onu, Meloni non sarà all’Assemblea generale a New York: al suo posto Tajani».",
    "Ministero degli Affari Esteri e della Cooperazione Internazionale, 20 settembre 2026 — «Tajani partecipa alla 81ª Assemblea generale delle Nazioni Unite».",
    "ANSA, 25 settembre 2026 — «Meloni, a Milano per vicinanza al settore moda, la sinistra non conosce l’Italia reale».",
    "Roll Call / Factbase, 2 ottobre 2026 — «Donald Trump Press Gaggle Before Marine One Departure».",
]
for source in sources:
    document.add_paragraph(source, style="List Bullet")

document.save(DOCX)
print(DOCX)
