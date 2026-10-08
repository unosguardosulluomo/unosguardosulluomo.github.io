"""Build the hybrid warfare dossier and its editorial placements."""

import email.utils
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
D = json.loads((ROOT / "editorial/guerra-ibrida.json").read_text(encoding="utf-8"))
SITE = "https://unosguardosulluomo.github.io/"
URL = SITE + D["slug"]


def esc(value):
    return html.escape(str(value), quote=True)


def rich_paragraph(paragraph):
    fragments = []
    for run in paragraph.runs:
        value = esc(run.text).replace("\n", "<br>")
        if not value:
            continue
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments) or esc(paragraph.text)


def normalize_published_text(value):
    replacements = {
        "nel SISMI — Servizio per le Informazioni e la Sicurezza Militare —": "nel Servizio per le Informazioni e la Sicurezza Militare (SISMI)",
        "del DIS — Dipartimento delle Informazioni per la Sicurezza —": "del Dipartimento delle Informazioni per la Sicurezza (DIS)",
        "La NATO considera": "L’Organizzazione del Trattato dell’Atlantico del Nord (NATO) considera",
        "L’Unione europea, nel luglio 2025": "L’Unione europea (UE), nel luglio 2025",
        "compreso il GRU — Direzione principale dello Stato maggiore delle Forze armate russe, cioè l’intelligence militare —": "compresa la Direzione principale dello Stato maggiore delle Forze armate russe (GRU), cioè l’intelligence militare,",
        "a APT28 — Advanced Persistent Threat 28 —,": "all’attore Advanced Persistent Threat 28 (APT28),",
        "con ORION 2026": "con l’Operazione di grande portata per eserciti resilienti, interoperabili, orientati al combattimento ad alta intensità e innovativi (ORION) 2026",
        "Il 7 ottobre Regno Unito e Germania": "L’8 ottobre Regno Unito e Germania",
        "Sala CSIRT — Computer Security Incident Response Team — mentre l’ACN — Agenzia per la Cybersicurezza Nazionale —": "Sala Computer Security Incident Response Team (CSIRT), mentre l’Agenzia per la Cybersicurezza Nazionale (ACN)",
        "del GNL — gas naturale liquefatto —": "del gas naturale liquefatto (GNL)",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value


def figure(image, width, height, alt):
    return (
        '<figure class="article-inline-image">'
        f'<img src="{esc(image)}" width="{width}" height="{height}" alt="{esc(alt)}" loading="lazy">'
        f'<figcaption>{esc(D["imageCaption"])}</figcaption></figure>'
    )


document = Document(ROOT / D["sourceDocx"])
body = []
contents = []
sources = []
editorial_note = ""
in_sources = False
heading_number = 0
caption_number = 0

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if not text:
        continue
    style = paragraph.style.name
    if style in {"Kicker", "Title", "Subtitle", "Deck"}:
        continue
    if style == "Caption":
        caption_number += 1
        if caption_number == 2:
            body.append(
                figure(
                    D["secondaryImage"],
                    D["secondaryImageWidth"],
                    D["secondaryImageHeight"],
                    D["secondaryImageAlt"],
                )
            )
        continue
    if style == "Heading 1" and text.casefold() == "fonti":
        in_sources = True
        continue
    if style == "Meta":
        editorial_note = normalize_published_text(rich_paragraph(paragraph))
        continue
    if in_sources:
        sources.append(text.removeprefix("•").strip())
        continue
    if style == "Heading 1":
        if text == "GPS e guerra elettronica":
            text = "Sistema di posizionamento globale (GPS) e guerra elettronica"
        heading_number += 1
        anchor = f"sezione-{heading_number}"
        contents.append(f'<li><a href="#{anchor}">{esc(text)}</a></li>')
        body.append(f'<h2 id="{anchor}">{esc(text)}</h2>')
    elif style == "Heading 2":
        if text == "GPS e guerra elettronica":
            text = "Sistema di posizionamento globale (GPS) e guerra elettronica"
        body.append(f"<h3>{esc(text)}</h3>")
    elif style == "Pullquote":
        body.append(f'<blockquote class="article-pullquote">{normalize_published_text(rich_paragraph(paragraph))}</blockquote>')
    else:
        body.append(f"<p>{normalize_published_text(rich_paragraph(paragraph))}</p>")

if len(sources) != 27:
    raise RuntimeError(f"Expected 27 documentary sources, found {len(sources)}")
if not editorial_note:
    raise RuntimeError("Editorial note not found")

word_count = len(re.findall(r"\b[\wÀ-ÖØ-öø-ÿ’'-]+\b", " ".join(p.text for p in document.paragraphs)))

organization = {
    "@type": ["Organization", "NewsMediaOrganization"],
    "@id": SITE + "#organization",
    "name": "Uno Sguardo sull’Uomo",
    "alternateName": ["UNO SGUARDO SULL'UOMO", "unosguardosulluomo"],
    "url": SITE,
    "description": "Testata editoriale indipendente italiana di approfondimento: indagini, dati, documenti e ricostruzione dei fatti oltre la narrazione.",
    "email": "mailto:unosguardosulluomo@gmail.com",
    "logo": {"@type": "ImageObject", "url": SITE + "assets/testata.webp"},
    "sameAs": [
        "https://www.facebook.com/profile.php?id=61593043131759",
        "https://www.linkedin.com/company/unosguardosulluomo/",
    ],
}
website = {
    "@type": "WebSite",
    "@id": SITE + "#website",
    "url": SITE,
    "name": "Uno Sguardo sull’Uomo",
    "alternateName": ["UNO SGUARDO SULL'UOMO", "unosguardosulluomo"],
    "description": "Dossier di geopolitica organizzati tra politica italiana, politica internazionale, economia e società.",
    "publisher": {"@id": SITE + "#organization"},
    "inLanguage": "it-IT",
}
schema = {
    "@context": "https://schema.org",
    "@graph": [
        organization,
        website,
        {
            "@type": "NewsArticle",
            "@id": URL + "#article",
            "headline": D["title"],
            "description": D["description"],
            "url": URL,
            "mainEntityOfPage": {"@type": "WebPage", "@id": URL},
            "datePublished": D["datePublished"],
            "dateModified": D["dateModified"],
            "image": [SITE + D["socialImage"]],
            "inLanguage": "it-IT",
            "isAccessibleForFree": True,
            "articleSection": D["category"],
            "genre": "Dossier geopolitico",
            "keywords": D["tags"],
            "about": [
                {"@type": "Thing", "name": "Geopolitica"},
                {"@type": "Thing", "name": D["category"]},
                *[{"@type": "Thing", "name": tag} for tag in D["tags"]],
            ],
            "wordCount": word_count,
            "author": {"@type": "Organization", "name": D["schemaAuthor"], "url": SITE + "chi-siamo.html"},
            "publisher": {"@id": SITE + "#organization"},
            "isPartOf": {"@id": SITE + "#website"},
        },
        {
            "@type": "BreadcrumbList",
            "itemListElement": [
                {"@type": "ListItem", "position": 1, "name": "Prima pagina", "item": SITE},
                {"@type": "ListItem", "position": 2, "name": D["category"], "item": SITE + D["categoryPath"]},
                {"@type": "ListItem", "position": 3, "name": D["title"], "item": URL},
            ],
        },
    ],
}

page = f'''<!doctype html>
<html lang="it">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="robots" content="index, follow, max-image-preview:large">
  <meta name="description" content="{esc(D['description'])}">
  <meta property="og:site_name" content="UNO SGUARDO SULL'UOMO">
  <meta property="og:title" content="{esc(D['title'])}">
  <meta property="og:description" content="{esc(D['description'])}">
  <meta property="og:type" content="article">
  <meta property="og:url" content="{URL}">
  <meta property="og:image" content="{SITE+D['socialImage']}">
  <meta property="og:image:alt" content="{esc(D['imageAlt'])}">
  <meta property="og:image:width" content="{D['socialImageWidth']}">
  <meta property="og:image:height" content="{D['socialImageHeight']}">
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="{esc(D['title'])}">
  <meta name="twitter:description" content="{esc(D['description'])}">
  <meta name="twitter:image" content="{SITE+D['socialImage']}">
  <meta property="article:published_time" content="{D['datePublished']}">
  <meta property="article:modified_time" content="{D['dateModified']}">
  <meta property="article:section" content="{esc(D['category'])}">
  <meta name="author" content="{esc(D['schemaAuthor'])}">
  {''.join(f'<meta property="article:tag" content="{esc(tag)}">' for tag in D['tags'])}
  <link rel="canonical" href="{URL}">
  <link rel="alternate" type="application/rss+xml" title="Uno Sguardo sull’Uomo" href="{SITE}feed.xml">
  <title>{esc(D['title'])} — Uno Sguardo sull'Uomo</title>
  <link rel="stylesheet" href="styles.css?v=20260924-mobile-1">
  <link rel="stylesheet" href="editorial-rules.css?v=20260812-1">
  <link rel="stylesheet" href="dossier-tables.css?v=20260925-1">
  <link rel="icon" type="image/svg+xml" href="assets/favicon.svg">
  <script type="application/ld+json" data-seo-schema>{json.dumps(schema, ensure_ascii=False, separators=(',', ':'))}</script>
</head>
<body><main class="newspaper">
  <div class="utility-bar"><span>Quotidiano indipendente di approfondimento</span><span>Politica, società e istituzioni</span><time class="current-date" datetime="{D['datePublished']}">{D['dateLabel']}</time></div>
  <header class="masthead"><a href="/" aria-label="Prima pagina"><img src="assets/testata.webp" alt="Uno Sguardo sull’Uomo"></a></header>
  <nav class="nav" aria-label="Navigazione principale"><a href="/">Prima pagina</a><a href="indagini.html">Indagini</a><a href="metodo.html">Metodo</a><a href="chi-siamo.html">Chi siamo</a><a href="contatti.html">Contatti</a></nav>
  <article class="article-page" data-macro-topic="Geopolitica" data-micro-topic="{esc(D['category'])}">
    <nav class="topic-breadcrumb" aria-label="Percorso tematico"><span class="topic-breadcrumb-macro">Geopolitica</span><span aria-hidden="true">›</span><a href="{D['categoryPath']}">{esc(D['category'])}</a><span aria-hidden="true">›</span><span aria-current="page">{esc(D['title'])}</span></nav>
    <header class="article-header"><a class="category-label" href="{D['categoryPath']}">{esc(D['category'])}</a><p class="eyebrow">DOSSIER</p><h1>{esc(D['title'])}</h1><p class="article-deck">{esc(D['deck'])}</p><div class="article-meta"><span>{esc(D['author'])}</span><time datetime="{D['datePublished']}">Pubblicato: {D['dateLabel']}</time><span>Tempo di lettura: {D['readingMinutes']} minuti</span></div><div class="topic-list" aria-label="Argomenti">{''.join('<span>'+esc(tag)+'</span>' for tag in D['tags'])}</div></header>
    <figure class="article-hero"><img src="{D['image']}" width="{D['imageWidth']}" height="{D['imageHeight']}" alt="{esc(D['imageAlt'])}" fetchpriority="high"><figcaption>{esc(D['imageCaption'])}</figcaption></figure>
    <div class="article-layout"><div class="article-body"><p class="article-lead">{esc(D['lead'])}</p><details class="dossier-contents"><summary>In questo dossier</summary><ol>{''.join(contents)}<li><a href="#fonti">Fonti</a></li><li><a href="#nota-editoriale">Nota editoriale</a></li></ol></details>
{chr(10).join(body)}
<section class="sources" aria-labelledby="fonti"><h2 id="fonti">FONTI</h2><ul class="sources-list">{''.join('<li>'+esc(source)+'</li>' for source in sources)}</ul></section>
<section class="method-note" aria-labelledby="nota-editoriale"><h2 id="nota-editoriale">NOTA EDITORIALE</h2><p>{editorial_note}</p></section>
    </div></div>
  </article>
  <footer class="footer"><span>Uno Sguardo sull’Uomo — {esc(D['category'])}</span><span><a href="privacy-cookie.html">Privacy e Cookie Policy</a></span><span><a href="mailto:unosguardosulluomo@gmail.com">unosguardosulluomo@gmail.com</a></span></footer>
</main><script src="date.js?v=20260826-audit-1"></script></body></html>
'''
(ROOT / D["slug"]).write_text(page, encoding="utf-8")


def card(kind="archive-card", heading="h3", eager=False):
    priority = ' fetchpriority="high"' if eager else ' loading="lazy"'
    return (
        f'<article class="{kind}" data-dossier="{esc(D["slug"])}">'
        f'<a class="story-image" href="{esc(D["slug"])}" aria-label="Leggi il dossier {esc(D["title"])}">'
        f'<img src="{esc(D["image"])}" width="{D["imageWidth"]}" height="{D["imageHeight"]}" alt="{esc(D["imageAlt"])}"{priority}></a>'
        f'<p class="eyebrow">{esc(D["cardEyebrow"])}</p>'
        f'<p class="archive-meta"><time datetime="{D["datePublished"]}">{D["dateLabel"]}</time></p>'
        f'<{heading}><a class="headline-link" href="{esc(D["slug"])}">{esc(D["title"])}</a></{heading}>'
        f'<p>{esc(D["description"])}</p>'
        f'<a class="read-more" href="{esc(D["slug"])}">Leggi il dossier →</a></article>'
    )


def dated_cards(markup):
    cards = re.findall(r'<article class="archive-card".*?</article>', markup, re.S)
    return sorted(cards, key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)


def change_card(markup, kind, heading):
    updated = re.sub(r'class="archive-card"', f'class="{kind}"', markup, count=1)
    updated = re.sub(r'<h3>', f'<{heading}>', updated, count=1)
    updated = re.sub(r'</h3>', f'</{heading}>', updated, count=1)
    updated = updated.replace(' fetchpriority="high"', ' loading="lazy"')
    return updated


def update_current_date(markup):
    return re.sub(
        r'<time class="current-date" datetime="[^"]+">.*?</time>',
        f'<time class="current-date" datetime="{D["datePublished"]}">{D["dateLabel"]}</time>',
        markup,
        count=1,
    )


def card_items(markup):
    items = []
    for position, article in enumerate(re.findall(r'<article class="[^"]*".*?</article>', markup, re.S), 1):
        href = re.search(r'<h[123]><a class="headline-link" href="([^"]+)">', article)
        name = re.search(r'<h[123]><a class="headline-link" href="[^"]+">(.*?)</a></h[123]>', article, re.S)
        if href and name:
            items.append({"@type": "ListItem", "position": position, "url": SITE + html.unescape(href.group(1)), "name": html.unescape(re.sub('<.*?>', '', name.group(1)))})
    return items


def rewrite_schema(markup, updater):
    match = re.search(r'<script type="application/ld\+json" data-seo-schema>(.*?)</script>', markup, re.S)
    if not match:
        raise RuntimeError("SEO schema not found")
    data = json.loads(match.group(1))
    updater(data)
    replacement = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    return markup[:match.start(1)] + replacement + markup[match.end(1):]


archive_path = ROOT / D["categoryPath"]
archive = archive_path.read_text(encoding="utf-8")
archive = re.sub(r'<article class="archive-card" data-dossier="' + re.escape(D["slug"]) + r'".*?</article>', "", archive, flags=re.S)
archive = re.sub(r'(<div class="archive-grid">)\s*', r'\1\n', archive, count=1)
archive = archive.replace('<div class="archive-grid">\n', '<div class="archive-grid">\n' + card() + '\n', 1)
archive = update_current_date(archive)


def update_archive_schema(data):
    collection = next(item for item in data["@graph"] if item.get("@type") == "CollectionPage")
    items = card_items(archive)
    collection["mainEntity"]["numberOfItems"] = len(items)
    collection["mainEntity"]["itemListElement"] = items


archive = rewrite_schema(archive, update_archive_schema)
archive_path.write_text(archive, encoding="utf-8")

category_cards = dated_cards(archive)
indagini_path = ROOT / "indagini.html"
indagini = indagini_path.read_text(encoding="utf-8")
section = re.compile(r'(<section class="archive-section"><div class="archive-section-header"><h2><a class="headline-link" href="archivio-politica-internazionale.html">Politica internazionale</a></h2></div><div class="archive-grid">).*?(</div><a class="category-archive-link" href="archivio-politica-internazionale.html">)', re.S)
indagini, count = section.subn(lambda match: match.group(1) + "".join(category_cards[:3]) + match.group(2), indagini, count=1)
if count != 1:
    raise RuntimeError("Politica internazionale selection not found")
indagini = update_current_date(indagini)


def update_indagini_schema(data):
    collection = next(item for item in data["@graph"] if item.get("@type") == "CollectionPage")
    items = card_items(indagini)
    collection["mainEntity"]["numberOfItems"] = len(items)
    collection["mainEntity"]["itemListElement"] = items


indagini = rewrite_schema(indagini, update_indagini_schema)
indagini_path.write_text(indagini, encoding="utf-8")

all_archive_cards = []
for category_archive in ("archivio-politica-italiana.html", "archivio-politica-internazionale.html", "archivio-economia.html", "archivio-societa.html"):
    all_archive_cards.extend(dated_cards((ROOT / category_archive).read_text(encoding="utf-8")))
all_archive_cards.sort(key=lambda item: re.search(r'datetime="(\d{4}-\d{2}-\d{2})"', item).group(1), reverse=True)

home_path = ROOT / "index.html"
home = home_path.read_text(encoding="utf-8")
lead = change_card(all_archive_cards[0], "lead-story", "h1").replace(' loading="lazy"', ' fetchpriority="high"', 1)
side = "".join(change_card(item, "side-story", "h2") for item in all_archive_cards[1:3])
latest = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[3:6])
other = "".join(change_card(item, "story-card", "h3") for item in all_archive_cards[6:9])
home, count = re.subn(r'<article class="lead-story".*?</article>', lambda _: lead, home, count=1, flags=re.S)
if count != 1:
    raise RuntimeError("Home lead not found")
home, count = re.subn(r'(<aside>).*?(</aside>)', lambda match: match.group(1) + side + match.group(2), home, count=1, flags=re.S)
if count != 1:
    raise RuntimeError("Home sidebar not found")
home, count = re.subn(r'(<section class="story-grid" aria-label="Ultime indagini">).*?(</section>)', lambda match: match.group(1) + latest + match.group(2), home, count=1, flags=re.S)
if count != 1:
    raise RuntimeError("Home latest grid not found")
home, count = re.subn(r'(<section class="story-grid" aria-label="Altre indagini recenti">).*?(</section>)', lambda match: match.group(1) + other + match.group(2), home, count=1, flags=re.S)
if count != 1:
    raise RuntimeError("Home other grid not found")
home = update_current_date(home)


def update_home_schema(data):
    webpage = next(item for item in data["@graph"] if item.get("@type") == "WebPage")
    webpage["name"] = D["title"]


home = rewrite_schema(home, update_home_schema)
home_path.write_text(home, encoding="utf-8")

article_entry = f'  <url><loc>{URL}</loc><lastmod>{D["dateModified"]}</lastmod><image:image><image:loc>{SITE + D["socialImage"]}</image:loc></image:image></url>\n'
sitemap_articles_path = ROOT / "sitemap-articles.xml"
sitemap_articles = sitemap_articles_path.read_text(encoding="utf-8")
sitemap_articles = re.sub(r'\s*<url><loc>' + re.escape(URL) + r'</loc>.*?</url>', "", sitemap_articles, flags=re.S)
sitemap_articles = re.sub(
    r'(xmlns:image="http://www.google.com/schemas/sitemap-image/1.1">)\s*',
    r'\1\n' + article_entry,
    sitemap_articles,
    count=1,
)
sitemap_articles_path.write_text(sitemap_articles, encoding="utf-8")

news_entry = f'  <url><loc>{URL}</loc><news:news><news:publication><news:name>Uno Sguardo sull’Uomo</news:name><news:language>it</news:language></news:publication><news:publication_date>{D["datePublished"]}</news:publication_date><news:title>{esc(D["title"])}</news:title></news:news></url>\n'
sitemap_news_path = ROOT / "sitemap-news.xml"
sitemap_news = sitemap_news_path.read_text(encoding="utf-8")
sitemap_news = re.sub(r'\s*<url><loc>' + re.escape(URL) + r'</loc>.*?</url>', "", sitemap_news, flags=re.S)
sitemap_news = re.sub(
    r'(xmlns:news="http://www.google.com/schemas/sitemap-news/0.9">)\s*',
    r'\1\n' + news_entry,
    sitemap_news,
    count=1,
)
sitemap_news_path.write_text(sitemap_news, encoding="utf-8")

pub_dt = datetime.fromisoformat(D["datePublished"]).replace(tzinfo=timezone.utc)
rss_date = email.utils.format_datetime(pub_dt)
feed_item = f'''    <item>
      <title>{esc(D['title'])}</title>
      <link>{URL}</link>
      <guid isPermaLink="true">{URL}</guid>
      <pubDate>{rss_date}</pubDate>
      <category>{esc(D['category'])}</category>
      <description>{esc(D['description'])}</description>
    </item>
'''
feed_path = ROOT / "feed.xml"
feed = feed_path.read_text(encoding="utf-8")
feed = re.sub(r'\s*<item>\s*<title>' + re.escape(esc(D["title"])) + r'</title>.*?</item>', "", feed, flags=re.S)
feed = feed.replace('    <atom:link href="https://unosguardosulluomo.github.io/feed.xml" rel="self" type="application/rss+xml"/>\n', '    <atom:link href="https://unosguardosulluomo.github.io/feed.xml" rel="self" type="application/rss+xml"/>\n' + feed_item, 1)
feed_path.write_text(feed, encoding="utf-8")

sitemap_pages_path = ROOT / "sitemap-pages.xml"
sitemap_pages = sitemap_pages_path.read_text(encoding="utf-8")
for page_url in (SITE, SITE + "indagini.html", SITE + D["categoryPath"]):
    sitemap_pages = re.sub(r'(<loc>' + re.escape(page_url) + r'</loc><lastmod>)[^<]+', r'\g<1>' + D["dateModified"], sitemap_pages)
sitemap_pages_path.write_text(sitemap_pages, encoding="utf-8")

for filename in ("sitemap.xml", "sitemap-google.xml", "sitemap-index.xml"):
    path = ROOT / filename
    markup = path.read_text(encoding="utf-8")
    markup = re.sub(r'<lastmod>[^<]+</lastmod>', f'<lastmod>{D["dateModified"]}</lastmod>', markup)
    path.write_text(markup, encoding="utf-8")

print("Built hybrid warfare dossier, editorial placements, feed and sitemaps")
